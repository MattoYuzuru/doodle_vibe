# -*- coding: utf-8 -*-
"""Объективные факты об оформлении отчёта студента — для «Проверки своего отчёта» (check-instruction.md).

    python report_check.py <отчёт.pdf | отчёт.docx> --out <папка>

Пишет в папку:
  facts.json   — измерения: страницы, поля, шрифты, интервалы, подписи рисунков и таблиц, ссылки на них в тексте,
                 заголовки, титульный лист, список источников, десятичные точки, кавычки, тире;
  report.txt   — текст отчёта по страницам (PDF) или по абзацам с пометками стилей (DOCX);
  pages\\pNN.png — страницы картинками (PDF; у DOCX — если на ПК есть Word и файл удалось перевести в PDF).

Скрипт ничего не оценивает и не исправляет — только измеряет. Оценку даёт агент проверки.
Коды выхода: 0 — готово; 2 — файл не открылся или это не PDF/DOCX.
"""
import argparse
import json
import os
import re
import statistics
import sys
import zipfile
from collections import Counter

PT_MM = 25.4 / 72
MAX_PAGES_PNG = 40
STRUCT = ["СОДЕРЖАНИЕ", "ОГЛАВЛЕНИЕ", "ВВЕДЕНИЕ", "ЦЕЛЬ РАБОТЫ", "ЦЕЛЬ", "ЗАДАНИЕ", "ХОД РАБОТЫ", "ХОД ВЫПОЛНЕНИЯ РАБОТЫ", "ТЕОРЕТИЧЕСКАЯ ЧАСТЬ",
          "ПРАКТИЧЕСКАЯ ЧАСТЬ", "РЕЗУЛЬТАТЫ", "ВЫВОД", "ВЫВОДЫ", "ЗАКЛЮЧЕНИЕ", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "СПИСОК ЛИТЕРАТУРЫ",
          "СПИСОК ИСТОЧНИКОВ", "СПИСОК ИСПОЛЬЗОВАННОЙ ЛИТЕРАТУРЫ", "БИБЛИОГРАФИЧЕСКИЙ СПИСОК", "ПРИЛОЖЕНИЕ", "КОНТРОЛЬНЫЕ ВОПРОСЫ", "ОТВЕТЫ НА КОНТРОЛЬНЫЕ ВОПРОСЫ"]
CAP_FIG = re.compile(r"^(Рисунок|Рис\.?)\s*(\d+(?:\.\d+)*)\s*([–—−\-.:]?)\s*(.*)$")
CAP_TAB = re.compile(r"^(Таблица|Табл\.?)\s*(\d+(?:\.\d+)*)\s*([–—−\-.:]?)\s*(.*)$")
REFS_HEAD = re.compile(r"^(список\s+(использованных\s+|использованной\s+)?(источников|литературы)|библиографический\s+список)\s*$", re.I)


def mm(pt):
    return round(pt * PT_MM, 1)


def top(counter, n=5):
    total = sum(counter.values()) or 1
    return [{"value": k, "share": round(v / total, 3)} for k, v in counter.most_common(n)]


def norm_font(name):
    name = re.sub(r"^[A-Z]{6}\+", "", name or "")
    return re.sub(r"[-,]?(Bold|Italic|Regular|MT|PS|Oblique|It|Bd)+$", "", re.sub(r"[-,](Bold|Italic|BoldItalic|Regular).*$", "", name)).strip() or name


def numbering(nums):
    """Порядок номеров подписей: пропуски, повторы, не по порядку."""
    plain = [int(n) for n in nums if n.isdigit()]
    out = {"numbers": nums, "duplicates": sorted({n for n in nums if nums.count(n) > 1}), "missing": [], "out_of_order": False}
    if plain:
        out["missing"] = [n for n in range(1, max(plain) + 1) if n not in plain]
        out["out_of_order"] = plain != sorted(plain)
    return out


def text_facts(pages_text):
    """Проверки по тексту: ссылки на рисунки и таблицы, источники, десятичные точки, кавычки, тире."""
    full = "\n".join(pages_text)
    out = {}
    out["decimal_point_numbers"] = {"count": 0, "examples": []}
    out["decimal_comma_numbers"] = len(re.findall(r"(?<![\d.,])\d+,\d+(?![\d,])", full))
    for i, t in enumerate(pages_text, 1):
        for line in t.split("\n"):
            # строки, похожие на код, команды и вывод программ, не считаем: там точка — норма
            if re.search(r"[{};=<>$#\\/|]|\b(def|import|return|print|docker|http|ms|GET|POST)\b", line):
                continue
            for m in re.finditer(r"(?<![\d.,\w])\d+\.\d+(?![\d.\w])", line):
                if not re.search(r"[А-Яа-яЁё]", line):
                    continue
                out["decimal_point_numbers"]["count"] += 1
                if len(out["decimal_point_numbers"]["examples"]) < 12:
                    out["decimal_point_numbers"]["examples"].append({"page": i, "number": m.group(0), "line": line.strip()[:90]})
    out["straight_quotes"] = len(re.findall(r'"[А-Яа-яЁё][^"\n]{0,80}"', full))
    out["guillemets"] = full.count("«")
    out["hyphen_as_dash"] = len(re.findall(r"[А-Яа-яЁё\d] - [А-Яа-яЁё\d]", full))
    out["en_dash"] = len(re.findall(r" – ", full))
    out["em_dash"] = len(re.findall(r" — ", full))
    out["double_spaces"] = len(re.findall(r"[А-Яа-яЁё,.;:]  +[А-Яа-яЁё]", full))
    # список источников и ссылки [N]
    lines = full.split("\n")
    start = next((k for k, l in enumerate(lines) if REFS_HEAD.match(l.strip())), None)
    cites = sorted({int(x) for m in re.finditer(r"\[(\d{1,3}(?:\s*[,;–-]\s*\d{1,3})*)\]", full) for x in re.findall(r"\d+", m.group(1))})
    refs = {"heading_found": start is not None, "entries": 0, "citations_in_text": cites, "never_cited": [], "cited_but_absent": [], "entries_preview": []}
    if start is not None:
        entries = []
        for l in lines[start + 1:]:
            s = l.strip()
            if re.match(r"^(приложение)\b", s, re.I):
                break
            m = re.match(r"^(\d{1,3})[.)]\s+(.{8,})", s)
            if m:
                entries.append((int(m.group(1)), m.group(2)))
        refs["entries"] = len(entries)
        refs["entries_preview"] = [f"{n}. {t[:110]}" for n, t in entries[:25]]
        nums = [n for n, _ in entries]
        refs["never_cited"] = [n for n in nums if n not in cites]
        refs["cited_but_absent"] = [n for n in cites if nums and n not in nums]
        refs["with_url"] = sum(1 for _, t in entries if re.search(r"https?://|URL", t))
        refs["with_access_date"] = sum(1 for _, t in entries if re.search(r"дата обращения", t, re.I))
    out["references"] = refs
    return out


def mentions(full, word_re, number, captions_count):
    """Сколько раз рисунок/таблица N упомянуты в тексте помимо самой подписи."""
    n = len(re.findall(rf"(?<![А-Яа-яЁё])({word_re})\s*{re.escape(number)}(?![\d.]\d)", full, re.I))
    return max(0, n - captions_count)


def title_facts(text):
    t = text[:2500]
    return {
        "text": t[:1500],
        "has_university": bool(re.search(r"МИСИС|университет|институт", t, re.I)),
        "has_work_kind": bool(re.search(r"отч[её]т|лабораторн|практическ|курсов|реферат|домашн", t, re.I)),
        "has_discipline_word": bool(re.search(r"дисциплин|по курсу|по предмету", t, re.I)),
        "has_group": bool(re.search(r"[А-ЯЁ]{2,6}-\d{2}(-[А-ЯЁ\d]+)*", t)),
        "has_done_by": bool(re.search(r"выполнил|студент", t, re.I)),
        "has_checked_by": bool(re.search(r"проверил|преподаватель|руководитель", t, re.I)),
        "has_city_year": bool(re.search(r"(Москва|г\.)\s*,?\s*20\d\d|20\d\d\s*(г\.?)?\s*$", t.strip(), re.I | re.M)),
        "has_variant": bool(re.search(r"вариант", t, re.I)),
    }


def is_struct(text):
    s = re.sub(r"[\s.:]+$", "", text.strip()).upper()
    s = re.sub(r"\s+[А-ЯA-Z]$", "", s) if s.startswith("ПРИЛОЖЕНИЕ") else s
    return s in STRUCT


MONO_FONTS = ("courier", "consol", "lucidacon", "cascadia", "firacode", "menlo", "monaco", "sourcecode", "terminal", "fixedsys", "mono")


def is_mono(name):
    """Моноширинный шрифт (листинги, код в строке)? По имени шрифта."""
    n = re.sub(r"[^a-z]", "", (name or "").lower()).replace("monotype", "")
    return any(k in n for k in MONO_FONTS)


def image_text_px(blob):
    """Размер картинки и высота строки текста на ней в пикселях (None — текста не видно).

    Строки текста — это полосы «чернил», разделённые полосами фона; берём медиану высот невысоких полос. У графика
    таких полос мало (подписи осей), у снимка терминала — десятки."""
    import io
    import numpy as np
    from PIL import Image
    im = Image.open(io.BytesIO(blob))
    w, h = im.size
    if w * h > 40_000_000:
        return w, h, None
    a = np.asarray(im.convert("L"), dtype=np.int16)
    ink = np.abs(a - int(np.bincount(a.ravel(), minlength=256).argmax())) > 60  # фон — самый частый тон
    rows = ink.sum(axis=1) > max(2, w * 0.004)
    runs, n = [], 0
    for r in list(rows) + [False]:
        if r:
            n += 1
        elif n:
            runs.append(n)
            n = 0
    small = [x for x in runs if 5 <= x <= 60]
    return w, h, (float(statistics.median(small)) if len(small) >= 3 else None)


# ---------- PDF ----------
def pdf_facts(path, out_dir, render=True, detail=False):
    import pymupdf
    doc = pymupdf.open(path)
    facts = {"pages": doc.page_count}
    fonts, sizes, ratios, indents = Counter(), Counter(), [], []
    left, right, topm, bottom = [], [], [], []
    justified = ragged = 0
    pages_text, captions_fig, captions_tab, headings, page_numbers, images, tables, empty = [], [], [], [], [], [], [], []
    det = {"lines": [], "fonts": Counter(), "sizes": Counter(), "page_sizes_mm": []}  # подробности для report_kit (detail=True)
    for pno, page in enumerate(doc, 1):
        W, H = page.rect.width, page.rect.height
        if pno == 1:
            facts["page_size_mm"] = [mm(W), mm(H)]
        det["page_sizes_mm"].append([mm(W), mm(H)])
        d = page.get_text("dict")
        lines_all = []
        img = sum(1 for b in d["blocks"] if b["type"] == 1 and (b["bbox"][2] - b["bbox"][0]) > 50 and (b["bbox"][3] - b["bbox"][1]) > 50)
        try:
            draw_img = len([i for i in page.get_image_info() if (i["bbox"][2] - i["bbox"][0]) > 50 and (i["bbox"][3] - i["bbox"][1]) > 50])
        except Exception:
            draw_img = img
        images.append(max(img, draw_img))
        try:
            tables.append(len(page.find_tables().tables))
        except Exception:
            tables.append(0)
        for b in d["blocks"]:
            if b["type"] != 0:
                continue
            blines = []
            for l in b["lines"]:
                spans = [s for s in l["spans"] if s["text"].strip()]
                if not spans:
                    continue
                text = "".join(s["text"] for s in l["spans"]).strip()
                x0, y0, x1, y1 = l["bbox"]
                size = max(spans, key=lambda s: len(s["text"]))["size"]
                bold = sum(len(s["text"]) for s in spans if s["flags"] & 16 or "Bold" in s["font"]) > len(text) / 2
                item = {"text": text, "x0": x0, "x1": x1, "y0": y0, "y1": y1, "size": size, "bold": bold, "base": spans[0]["origin"][1]}
                is_pn = bool(re.fullmatch(r"[-–— ]*\d{1,4}[-–— ]*", text) and (y0 > H * 0.9 or y1 < H * 0.08))
                if detail:  # строка целиком и её часть без моноширинных кусков (код в строке), шрифты основного текста
                    plain = [s for s in spans if not is_mono(s["font"])]
                    det["lines"].append({"page": pno, "text": text, "plain": "".join(s["text"] for s in plain).strip(), "y": round(y0, 1), "page_number": is_pn,
                                         "mono": sum(len(s["text"]) for s in plain) < len("".join(s["text"] for s in spans)) / 2})
                    for s in ([] if is_pn else plain):
                        det["fonts"][norm_font(s["font"])] += len(s["text"])
                        det["sizes"][round(s["size"] * 2) / 2] += len(s["text"])
                # номер страницы: одиночное число в верхних 8 % или нижних 10 % листа
                if re.fullmatch(r"[-–— ]*\d{1,4}[-–— ]*", text) and (y0 > H * 0.9 or y1 < H * 0.08):
                    cx = (x0 + x1) / 2
                    page_numbers.append({"page": pno, "number": int(re.sub(r"\D", "", text)), "where": ("низ" if y0 > H / 2 else "верх") + (", центр" if abs(cx - W / 2) < W * 0.1 else ", справа" if cx > W / 2 else ", слева")})
                    continue
                for s in spans:
                    fonts[norm_font(s["font"])] += len(s["text"])
                    sizes[round(s["size"] * 2) / 2] += len(s["text"])
                blines.append(item)
                lines_all.append(item)
            for k, l in enumerate(blines):
                t = l["text"]
                centered = abs((l["x0"] + l["x1"]) / 2 - W / 2) < W * 0.06
                mf, mt = CAP_FIG.match(t), CAP_TAB.match(t)
                if mf and k == 0 and (mf.group(3) or not mf.group(4) or centered):
                    captions_fig.append({"page": pno, "word": mf.group(1), "number": mf.group(2), "separator": mf.group(3) or "", "text": mf.group(4)[:100], "centered": centered, "ends_with_dot": t.rstrip().endswith("."), "size": round(l["size"], 1)})
                elif mt and k == 0 and (mt.group(3) or not mt.group(4) or len(t) < 120):
                    captions_tab.append({"page": pno, "word": mt.group(1), "number": mt.group(2), "separator": mt.group(3) or "", "text": mt.group(4)[:100], "centered": centered, "at_left": l["x0"] < W * 0.25, "ends_with_dot": t.rstrip().endswith("."), "size": round(l["size"], 1)})
                elif len(t) < 110 and (is_struct(t) or (re.match(r"^\d+(\.\d+)*\.?\s+[А-ЯЁA-Z]", t) and not t.rstrip().endswith((".", ",", ";", ":")))):
                    headings.append({"page": pno, "text": t, "bold": l["bold"], "size": round(l["size"], 1), "centered": centered, "uppercase": t.upper() == t, "first_on_page": len(lines_all) - len(blines) + k == 0, "ends_with_dot": t.rstrip().endswith("."), "structural": is_struct(t)})
        text = page.get_text("text")
        pages_text.append(text)
        if not text.strip() and not images[-1]:
            empty.append(pno)
        if lines_all and pno > 1:  # титульный лист в поля не идёт: он свёрстан по-своему
            wide = [l for l in lines_all if (l["x1"] - l["x0"]) > W * 0.5] or lines_all
            pl, pr = min(l["x0"] for l in lines_all), max(l["x1"] for l in wide)
            left.append(pl); right.append(W - pr)
            topm.append(min(l["y0"] for l in lines_all)); bottom.append(H - max(l["y1"] for l in lines_all))
            # Интервал, красная строка и выравнивание — по соседним строкам основного текста. По блокам PyMuPDF нельзя:
            # при полуторном интервале он каждую строку считает отдельным блоком.
            rows = sorted(lines_all, key=lambda l: (round(l["base"]), l["x0"]))
            for r1, r2 in zip(rows, rows[1:]):
                step = r2["base"] - r1["base"]
                if abs(r1["size"] - r2["size"]) > 0.6 or not (r1["size"] * 0.9 < step < r1["size"] * 3) or (r1["x1"] - r1["x0"]) < (pr - pl) * 0.6:
                    continue
                if r2["x0"] - pl < 3:  # следующая строка начинается от поля — это продолжение того же абзаца
                    ratios.append(step / r1["size"])
                    if pr - r1["x1"] < 3:
                        justified += 1
                    else:
                        ragged += 1
                    indents.append(r1["x0"] - pl if 5 < r1["x0"] - pl < 70 else 0)
    med = lambda v: mm(statistics.median(v)) if v else None
    low = lambda v: mm(sorted(v)[max(0, len(v) // 10)]) if v else None
    body_size = sizes.most_common(1)[0][0] if sizes else None
    ratio = statistics.median(ratios) if ratios else None
    facts["margins_mm_estimate"] = {"left": low(left), "right": low(right), "top": low(topm), "bottom_text": low(bottom),
                                    "note": "оценка по границам текста на страницах после титульной: левое и верхнее — точные, правое — по самым длинным строкам, нижнее — до последней строки текста (поле не больше этого)"}
    facts["fonts"] = top(fonts)
    facts["font_sizes_pt"] = top(sizes)
    facts["body_font_size_pt"] = body_size
    facts["line_step_to_font_size"] = round(ratio, 2) if ratio else None
    facts["line_spacing_estimate"] = (round(ratio / 1.15, 2) if ratio else None)
    facts["line_spacing_note"] = "оценка: шаг строк, делённый на кегль и на 1,15 (у Times New Roman одинарный интервал ≈ 1,15 кегля); 1,5 — полуторный"
    facts["first_line_indent_mm"] = med([i for i in indents if i > 2]) if [i for i in indents if i > 2] else None
    facts["first_line_indent_note"] = "красная строка: отступ первых строк абзацев от левого поля; null — абзацев с отступом не нашлось"
    facts["justified_lines_share"] = round(justified / (justified + ragged), 2) if justified + ragged else None
    facts["justified_note"] = "доля полных строк основного текста, доходящих до правого поля: около 1 — выравнивание по ширине, заметно меньше — по левому краю"
    nums = {p["page"]: p["number"] for p in page_numbers}
    facts["page_numbers"] = {"pages_with_number": len(nums), "positions": top(Counter(p["where"] for p in page_numbers), 3), "on_title_page": 1 in nums,
                             "pages_without_number": [p for p in range(2, doc.page_count + 1) if p not in nums][:30],
                             "mismatch": [{"page": p, "printed": n} for p, n in nums.items() if p != n][:10]}
    full = "\n".join(pages_text)
    for c in captions_fig:
        c["mentions_in_text"] = mentions(full, r"рисун\w*|рис\.", c["number"], sum(1 for x in captions_fig if x["number"] == c["number"]))
    for c in captions_tab:
        c["mentions_in_text"] = mentions(full, r"таблиц\w*|табл\.", c["number"], sum(1 for x in captions_tab if x["number"] == c["number"]))
    facts["figures"] = {"captions": captions_fig, "numbering": numbering([c["number"] for c in captions_fig]), "images_total": sum(images), "images_by_page": {str(i + 1): n for i, n in enumerate(images) if n}}
    facts["tables"] = {"captions": captions_tab, "numbering": numbering([c["number"] for c in captions_tab]), "tables_found_total": sum(tables), "tables_by_page": {str(i + 1): n for i, n in enumerate(tables) if n}}
    # нумерованные строки без выделения — это пункты списка, а не заголовки
    facts["headings"] = [h for h in headings if h["structural"] or h["bold"] or (body_size and h["size"] > body_size + 0.4)][:80]
    facts["empty_pages"] = empty
    facts["title_page"] = title_facts(pages_text[0] if pages_text else "")
    facts["words"] = len(re.findall(r"[А-Яа-яЁёA-Za-z]{2,}", full))
    facts.update(text_facts(pages_text))
    if detail:
        facts["detail"] = {"lines": det["lines"], "fonts": top(det["fonts"]), "sizes": top(det["sizes"]), "page_sizes_mm": det["page_sizes_mm"]}
    with open(os.path.join(out_dir, "report.txt"), "w", encoding="utf-8") as f:
        for i, t in enumerate(pages_text, 1):
            f.write(f"=== Страница {i} ===\n{t.strip()}\n\n")
    if render:
        pd = os.path.join(out_dir, "pages")
        os.makedirs(pd, exist_ok=True)
        for i, page in enumerate(doc, 1):
            if i > MAX_PAGES_PNG:
                break
            page.get_pixmap(dpi=90).save(os.path.join(pd, f"p{i:02d}.png"))
        facts["pages_png"] = min(doc.page_count, MAX_PAGES_PNG)
    return facts


# ---------- DOCX ----------
def docx_facts(path, out_dir, detail=False):
    import docx
    from docx.shared import Length
    d = docx.Document(path)
    facts = {}
    s0 = d.sections[0]
    L = lambda v: round(Length(v).mm, 1) if v is not None else None
    facts["page_size_mm"] = [L(s0.page_width), L(s0.page_height)]
    facts["margins_mm"] = {"left": L(s0.left_margin), "right": L(s0.right_margin), "top": L(s0.top_margin), "bottom": L(s0.bottom_margin)}
    if len(d.sections) > 1:
        facts["sections"] = [{"left": L(s.left_margin), "right": L(s.right_margin), "top": L(s.top_margin), "bottom": L(s.bottom_margin), "landscape": bool(s.orientation)} for s in d.sections]

    def chain(style):
        while style is not None:
            yield style
            style = style.base_style

    def eff_font(run, par, attr):
        v = getattr(run.font, attr)
        if v is None and run.style is not None:
            v = next((getattr(s.font, attr) for s in chain(run.style) if getattr(s.font, attr) is not None), None)
        if v is None:
            v = next((getattr(s.font, attr) for s in chain(par.style) if getattr(s.font, attr) is not None), None)
        return v

    # Word после сохранения пишет выравнивание словами «start»/«end» (вместо left/right) — python-docx таких не знает
    # и падает (5 окт 2026: проверка отчёта остановилась на «has no XML mapping for 'end'»). Читаем такое значение сами.
    JC = {"start": 0, "left": 0, "center": 1, "end": 2, "right": 2, "both": 3, "distribute": 3}

    def fmt_attr(fmt, attr):
        try:
            return getattr(fmt, attr)
        except Exception:
            if attr != "alignment":
                raise
            jc = fmt._element.find(".//{*}jc") if getattr(fmt, "_element", None) is not None else None
            val = jc.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") if jc is not None else None
            return JC.get(val)

    def eff_par(par, attr):
        v = fmt_attr(par.paragraph_format, attr)
        if v is None:
            v = next((fmt_attr(s.paragraph_format, attr) for s in chain(par.style) if fmt_attr(s.paragraph_format, attr) is not None), None)
        return v

    # шрифт и кегль по умолчанию из docDefaults (когда ни в прогоне, ни в стиле не задано)
    dflt_font, dflt_size = None, None
    try:
        rpr = d.styles.element.find(".//{*}docDefaults/{*}rPrDefault/{*}rPr")
        if rpr is not None:
            rf, sz = rpr.find("{*}rFonts"), rpr.find("{*}sz")
            if rf is not None:
                dflt_font = rf.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii")
            if sz is not None:
                dflt_size = int(sz.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")) / 2
    except Exception:
        pass
    ALIGN = {None: "по левому краю (по умолчанию)", 0: "по левому краю", 1: "по центру", 2: "по правому краю", 3: "по ширине"}
    fonts, sizes, spacing, indent, align = Counter(), Counter(), Counter(), Counter(), Counter()
    captions_fig, captions_tab, headings, lines, order = [], [], [], [], []
    body = d.element.body
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    for child in body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(child, d)
            t = p.text.strip()
            has_img = bool(child.findall(".//{*}drawing") or child.findall(".//{*}pict"))
            order.append({"kind": "image" if has_img and not t else "p", "text": t})
            if has_img:
                lines.append("[рисунок]")
            if not t:
                continue
            style = p.style.name if p.style is not None else ""
            al = ALIGN.get(int(eff_par(p, "alignment")) if eff_par(p, "alignment") is not None else None, "другое")
            for r in p.runs:
                n = len(r.text)
                if not n:
                    continue
                fonts[eff_font(r, p, "name") or dflt_font or "не задан (шрифт темы)"] += n
                sz = eff_font(r, p, "size")
                sizes[round(sz.pt * 2) / 2 if sz is not None else dflt_size or "не задан"] += n
            mf, mt = CAP_FIG.match(t), CAP_TAB.match(t)
            is_head = style.lower().startswith(("heading", "заголовок", "title")) or (len(t) < 110 and is_struct(t))
            if mf and (mf.group(3) or not mf.group(4) or len(t) < 150):
                captions_fig.append({"word": mf.group(1), "number": mf.group(2), "separator": mf.group(3) or "", "text": mf.group(4)[:100], "align": al, "ends_with_dot": t.endswith("."), "index": len(order) - 1})
            elif mt and (mt.group(3) or not mt.group(4) or len(t) < 150):
                captions_tab.append({"word": mt.group(1), "number": mt.group(2), "separator": mt.group(3) or "", "text": mt.group(4)[:100], "align": al, "ends_with_dot": t.endswith("."), "index": len(order) - 1})
            elif is_head:
                bold = bool(p.runs) and all((r.bold if r.bold is not None else eff_font(r, p, "bold")) for r in p.runs if r.text.strip())
                headings.append({"text": t[:110], "style": style, "bold": bold, "align": al, "uppercase": t.upper() == t, "ends_with_dot": t.endswith("."), "page_break_before": bool(eff_par(p, "page_break_before"))})
            elif len(t) > 150:  # основной текст: интервалы, отступы и выравнивание считаем по настоящим абзацам
                ls = eff_par(p, "line_spacing")
                spacing[round(float(ls), 2) if isinstance(ls, float) else (f"{ls.pt} пт" if ls is not None else "одинарный (по умолчанию)")] += len(t)
                fi = eff_par(p, "first_line_indent")
                indent[round(fi.mm / 2.5) * 2.5 if fi is not None else 0] += len(t)
                align[al] += len(t)
            lines.append((f"[{style}] " if is_head else "") + t)
        elif tag == "tbl":
            tb = Table(child, d)
            order.append({"kind": "table", "text": ""})
            lines.append("[таблица]")
            for row in tb.rows[:60]:
                seen, cells = set(), []
                for c in row.cells:
                    if id(c._tc) not in seen:
                        seen.add(id(c._tc)); cells.append(c.text.strip().replace("\n", " "))
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("[конец таблицы]")
    text = "\n".join(lines)
    # где подпись относительно таблицы и рисунка
    nxt = lambda i: next((o for o in order[i + 1:] if o["kind"] != "p" or o["text"]), None)
    prv = lambda i: next((o for o in reversed(order[:i]) if o["kind"] != "p" or o["text"]), None)
    for c in captions_tab:
        i = c.pop("index")
        c["position"] = "над таблицей" if (nxt(i) or {}).get("kind") == "table" else "под таблицей" if (prv(i) or {}).get("kind") == "table" else "рядом таблицы нет"
        c["mentions_in_text"] = mentions(text, r"таблиц\w*|табл\.", c["number"], sum(1 for x in captions_tab if x["number"] == c["number"]))
    for c in captions_fig:
        i = c.pop("index")
        c["position"] = "под рисунком" if (prv(i) or {}).get("kind") == "image" else "над рисунком" if (nxt(i) or {}).get("kind") == "image" else "рисунок не в соседнем абзаце"
        c["mentions_in_text"] = mentions(text, r"рисун\w*|рис\.", c["number"], sum(1 for x in captions_fig if x["number"] == c["number"]))
    n_img_all = len(body.findall(".//{*}drawing")) + len(body.findall(".//{*}pict"))
    facts["fonts"] = top(fonts)
    facts["font_sizes_pt"] = top(sizes)
    facts["body_line_spacing"] = top(spacing, 4)
    facts["body_first_line_indent_mm"] = top(indent, 4)
    facts["body_alignment"] = top(align, 4)
    facts["figures"] = {"captions": captions_fig, "numbering": numbering([c["number"] for c in captions_fig]), "images_total": n_img_all}
    facts["tables"] = {"captions": captions_tab, "numbering": numbering([c["number"] for c in captions_tab]), "tables_found_total": len(d.tables)}
    facts["headings"] = headings[:80]
    xml = body.xml
    facts["has_auto_toc"] = "TOC \\" in xml or 'w:val="Table of Contents"' in xml
    foot = "".join((s.footer._element.xml if s.footer is not None else "") + (s.header._element.xml if s.header is not None else "") for s in d.sections)
    facts["page_number_field"] = bool(re.search(r"PAGE\b", foot))
    facts["different_first_page"] = bool(s0.different_first_page_header_footer)
    try:
        with zipfile.ZipFile(path) as z:
            m = re.search(r"<Pages>(\d+)</Pages>", z.read("docProps/app.xml").decode("utf-8", "ignore"))
            facts["pages_in_file_properties"] = int(m.group(1)) if m else None
    except Exception:
        facts["pages_in_file_properties"] = None
    first = "\n".join(o["text"] for o in order[:45] if o["text"])
    facts["title_page"] = title_facts(first)
    facts["words"] = len(re.findall(r"[А-Яа-яЁёA-Za-z]{2,}", text))
    facts.update(text_facts([text]))
    for ex in facts["decimal_point_numbers"]["examples"]:
        ex.pop("page", None)
    if detail:
        facts["detail"] = docx_detail(d, eff_font, eff_par, dflt_font, dflt_size)
    with open(os.path.join(out_dir, "report.txt"), "w", encoding="utf-8") as f:
        f.write(text + "\n")
    return facts


TITLE_SCAN = 80  # титульную часть ищем среди первых стольких элементов тела


def docx_detail(d, eff_font, eff_par, dflt_font, dflt_size):
    """Подробные измерения для сборщика отчётов (report_kit check): каждый абзац, раздел, таблица и рисунок.

    В facts.json не попадают (docx_facts(..., detail=True) кладёт их в ключ detail только по просьбе)."""
    from docx.shared import Length
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    body = d.element.body
    kids = list(body.iterchildren())
    tag = lambda el: el.tag.split("}")[1]

    def style_of(p):
        try:
            st = p.style
            return (st.style_id or "", (st.name or "").lower()) if st is not None else ("", "")
        except Exception:
            return "", ""

    def par(p, attr):
        try:
            return eff_par(p, attr)
        except Exception:  # значение, которого python-docx не знает
            return None

    # Граница титульной части: разрыв страницы или раздела (включительно), абзац «с новой страницы», заголовок,
    # абзац сборщика (стиль RK…), «СОДЕРЖАНИЕ/ОГЛАВЛЕНИЕ/ВВЕДЕНИЕ» (не включая).
    end, title_section = None, False
    for k, child in enumerate(kids[:TITLE_SCAN]):
        if tag(child) != "p":
            continue
        p = Paragraph(child, d)
        sid, name = style_of(p)
        word = re.sub(r"[\s.:]+$", "", p.text.strip()).upper()
        if sid.startswith("RK") or name.startswith(("heading", "заголовок")) or par(p, "page_break_before") or word in ("СОДЕРЖАНИЕ", "ОГЛАВЛЕНИЕ", "ВВЕДЕНИЕ"):
            end = k
            break
        if child.find(f"{W}pPr/{W}sectPr") is not None:
            end, title_section = k + 1, True
            break
        if any(br.get(W + "type") == "page" for br in child.iter(W + "br")):
            end = k + 1
            break
    if end is None:  # границы нет: короткий документ — титул целиком, длинный — без титульной части
        end = len(kids) if len(kids) <= TITLE_SCAN else 0

    # межстрочный интервал по умолчанию из docDefaults (когда ни в абзаце, ни в стилях не задан)
    dflt_spacing = 1.0
    sp = d.styles.element.find(f".//{W}docDefaults/{W}pPrDefault/{W}pPr/{W}spacing")
    if sp is not None and sp.get(W + "line") and sp.get(W + "lineRule", "auto") == "auto":
        dflt_spacing = round(int(sp.get(W + "line")) / 240, 2)
    ALIGN = {None: "left", 0: "left", 1: "center", 2: "right", 3: "justify"}
    pars, tables = [], []
    for k, child in enumerate(kids):
        kind = tag(child)
        if kind not in ("p", "tbl"):
            continue
        for pel in child.iter(W + "p"):
            p = Paragraph(pel, d)
            text, plain, fonts, sizes = [], [], Counter(), Counter()
            for rel in pel.iter(W + "r"):
                r = Run(rel, p)
                s = r.text
                if not s:
                    continue
                try:
                    font, sz = eff_font(r, p, "name") or dflt_font, eff_font(r, p, "size")
                except Exception:
                    font, sz = dflt_font, None
                text.append(s)
                fonts[font or ""] += len(s)
                sizes[round(sz.pt * 2) / 2 if sz is not None else dflt_size or 0] += len(s)
                if not is_mono(font):
                    plain.append(s)
            text = "".join(text)
            if not text.strip():
                continue
            sid, name = style_of(p)
            ls, fi, al = par(p, "line_spacing"), par(p, "first_line_indent"), par(p, "alignment")
            pars.append({"k": k, "text": text, "plain": "".join(plain), "style": sid, "style_name": name, "table": kind == "tbl", "title": k < end,
                         "mono": len("".join(plain)) < len(text) / 2, "fonts": dict(fonts), "sizes": dict(sizes),
                         "spacing": round(float(ls), 2) if isinstance(ls, float) else (f"{ls.pt} пт" if ls is not None else dflt_spacing),
                         "indent_mm": round(fi.mm, 1) if fi is not None else 0.0, "align": ALIGN.get(int(al) if al is not None else None, "other")})
        if kind == "tbl":
            cap = child.find(f"{W}tblPr/{W}tblCaption")
            t = {"k": k, "title": k < end, "caption": cap.get(W + "val") if cap is not None else "", "cells": 0, "empty": 0, "decimal_point": 0, "decimal_point_examples": []}
            for tc in child.iter(W + "tc"):
                vm = tc.find(f"{W}tcPr/{W}vMerge")
                if vm is not None and vm.get(W + "val") != "restart":
                    continue  # продолжение объединённой ячейки
                s = "".join(x.text or "" for x in tc.iter(W + "t")).strip()
                t["cells"] += 1
                if not s and not any(tag(x) in ("drawing", "pict", "oMath") for x in tc.iter() if isinstance(x.tag, str)):
                    t["empty"] += 1
                if re.fullmatch(r"-?\d+\.\d+", s):
                    t["decimal_point"] += 1
                    if len(t["decimal_point_examples"]) < 5:
                        t["decimal_point_examples"].append(s)
            tables.append(t)
    L = lambda v: round(Length(v).mm, 1) if v is not None else None
    sections = [{"page_mm": [L(s.page_width), L(s.page_height)], "left": L(s.left_margin), "right": L(s.right_margin), "top": L(s.top_margin), "bottom": L(s.bottom_margin),
                 "landscape": bool(s.orientation), "title": title_section and i == 0} for i, s in enumerate(d.sections)]
    images = []
    for dr in list(body.iter(W + "drawing"))[:80]:
        ext, blip = dr.find(".//{*}extent"), dr.find(".//{*}blip")
        try:
            w, h, tp = image_text_px(d.part.related_parts[blip.get(R + "embed")].blob)
            images.append({"width_mm": round(int(ext.get("cx")) / 36000, 1), "px": [w, h], "text_px": tp})
        except Exception:
            continue  # связанный файл, векторная картинка, фигура без изображения
    return {"title_end": end, "sections": sections, "pars": pars, "tables": tables, "images": images}


def docx_to_pdf(path, out_pdf):
    """DOCX → PDF через Word, если он есть на ПК (для снимков страниц и числа страниц). Нет Word — возвращает False."""
    try:
        import pythoncom
        import win32com.client
    except Exception:
        return False
    word = None
    try:
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(os.path.abspath(path), ReadOnly=True, AddToRecentFiles=False, ConfirmConversions=False)
        doc.ExportAsFixedFormat(os.path.abspath(out_pdf), 17)  # 17 — PDF
        doc.Close(False)
        return os.path.exists(out_pdf)
    except Exception:
        return False
    finally:
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--out", required=True)
    ap.add_argument("--no-word", action="store_true", help="не переводить DOCX в PDF через Word")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    ext = os.path.splitext(a.file)[1].lower()
    try:
        if ext == ".pdf":
            facts = {"kind": "pdf", **pdf_facts(a.file, a.out)}
        elif ext == ".docx":
            facts = {"kind": "docx", **docx_facts(a.file, a.out)}
            pdf = os.path.join(a.out, "report-from-docx.pdf")
            if not a.no_word and docx_to_pdf(a.file, pdf):
                layout = pdf_facts(pdf, os.path.join(a.out, "_pdf"), render=False) if os.makedirs(os.path.join(a.out, "_pdf"), exist_ok=True) is None else {}
                import pymupdf
                doc = pymupdf.open(pdf)
                pd = os.path.join(a.out, "pages")
                os.makedirs(pd, exist_ok=True)
                for i, page in enumerate(doc, 1):
                    if i > MAX_PAGES_PNG:
                        break
                    page.get_pixmap(dpi=90).save(os.path.join(pd, f"p{i:02d}.png"))
                facts["pages"] = doc.page_count
                facts["pages_png"] = min(doc.page_count, MAX_PAGES_PNG)
                facts["page_numbers"] = layout.get("page_numbers")
                facts["empty_pages"] = layout.get("empty_pages")
                facts["layout_from_pdf"] = {k: layout.get(k) for k in ("body_font_size_pt", "line_spacing_estimate", "first_line_indent_mm", "justified_lines_share")}
            else:
                facts["pages"] = facts.get("pages_in_file_properties")
                facts["pages_note"] = "Word на ПК не найден: число страниц — из свойств файла (может быть неточным), снимков страниц нет"
        else:
            print("Нужен файл PDF или DOCX", file=sys.stderr)
            return 2
    except Exception as e:
        print(f"Файл не открылся: {e}", file=sys.stderr)
        return 2
    facts["file"] = os.path.basename(a.file)
    with open(os.path.join(a.out, "facts.json"), "w", encoding="utf-8") as f:
        json.dump(facts, f, ensure_ascii=False, indent=1)
    print(f"готово: {facts.get('pages')} стр., рисунков с подписью {len(facts['figures']['captions'])}, таблиц с подписью {len(facts['tables']['captions'])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
