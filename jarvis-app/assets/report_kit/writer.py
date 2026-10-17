# Запись DOCX из промежуточного представления (список блоков от mdparse): свои стили RK*, лист и поля, номер
# страницы, нумерация заголовков/рисунков/таблиц/листингов/формул, подстановка номеров в ссылки, титул, завершение
# (одинаковый вход → побайтно одинаковый файл) и самопроверка собранного файла.
import copy
import csv
import datetime
import io
import os
import posixpath
import re
import zipfile

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Pt, Twips
from docx.text.paragraph import Paragraph
from lxml import etree

from . import title as title_mod
from .common import (EQ_TABLE_MARK, EXIT_SELF, EXIT_WORD, STYLE_BODY, STYLE_CAPTION_FIG, STYLE_CAPTION_LST, STYLE_CAPTION_TAB,
                     STYLE_FORMULA, STYLE_LIST, STYLE_LISTING, STYLE_TABLE_TEXT, KitError, MathError, clean_text, resolve_path)

STYLE_TITLE = title_mod.STYLE_TITLE
PAGE_SIZES = {"A4": (11906, 16838), "A3": (16838, 23811), "A5": (8391, 11906), "LETTER": (12240, 15840)}  # твипы
JC = {"justify": "both", "left": "left", "center": "center", "right": "right"}
ALIGN = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT,
         "left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
         "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
IMAGE_EXT = (".png", ".jpg", ".jpeg", ".gif", ".bmp")
FIXED_DATE = datetime.datetime(2000, 1, 1)
EMU_TWIP, EMU_CM = 635, 360000
ROW_GAP = int(0.3 * EMU_CM)  # зазор между рисунками в ряд
TOC_STUB = "Обновите поле оглавления: выделить всё (Ctrl+A) и нажать F9"  # заменит Word при конвертации (topdf)
LEFTOVERS = ("{#", "@fig:", "@tbl:", "@eq:", "@lst:", "$$", "[TOC]", "![")  # следы разметки, которых в отчёте быть не должно
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_DECIMAL = re.compile(r"^-?\d+\.\d+$")
SOFT = chr(0xAD)  # место возможного переноса внутри _soft_wrap; в документ уходит элементом w:softHyphen


def _tw_mm(mm):
    return int(round(float(mm) * 1440 / 25.4))


def _tw_cm(cm):
    return int(round(float(cm) * 1440 / 2.54))


def page_metrics(style):
    """Лист и поля из профиля в твипах: (ширина, высота, {left, right, top, bottom}, ширина текста, высота текста)."""
    page = style.get("page") or {}
    name = str(page.get("size", "A4")).upper()
    if name not in PAGE_SIZES:
        raise KitError(f"style.json: неизвестный размер листа page.size = {page.get('size')!r}", hint="допустимо: A4, A3, A5, Letter")
    w, h = PAGE_SIZES[name]
    if page.get("orientation") == "landscape":
        w, h = h, w
    mar = {k: _tw_mm((page.get("margins_mm") or {}).get(k, d)) for k, d in (("left", 30), ("right", 15), ("top", 20), ("bottom", 20))}
    return w, h, mar, w - mar["left"] - mar["right"], h - mar["top"] - mar["bottom"]


# ---------- стили ----------

def _style_xml(sid, name, font, size_pt, lang, *, align="left", first=0, left=0, line=240, before=0, after=0, bold=False,
               caps=False, keep_next=False, keep_lines=False, page_break=False, outline=None, next_id=None, ui=None,
               tab_right=None, custom=True):
    """Стиль абзаца со всеми свойствами явно и без basedOn: от Normal и темы основы он не зависит."""
    sz = int(round(float(size_pt) * 2))
    ppr = ""
    ppr += "<w:keepNext/>" if keep_next else ""
    ppr += "<w:keepLines/>" if keep_lines else ""
    ppr += "<w:pageBreakBefore/>" if page_break else ""
    ppr += "<w:widowControl/>"
    ppr += f'<w:tabs><w:tab w:val="right" w:leader="dot" w:pos="{tab_right}"/></w:tabs>' if tab_right else ""
    ppr += "<w:suppressAutoHyphens/>" if outline is not None else ""
    ppr += f'<w:spacing w:before="{before}" w:after="{after}" w:line="{line}" w:lineRule="auto"/>'
    ppr += f'<w:ind w:left="{left}" w:right="0" w:firstLine="{first}"/>'
    ppr += f'<w:jc w:val="{JC.get(align, "left")}"/>'
    ppr += f'<w:outlineLvl w:val="{outline}"/>' if outline is not None else ""
    rpr = f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}" w:eastAsia="{font}"/>'
    rpr += f'<w:b w:val="{int(bool(bold))}"/><w:bCs w:val="{int(bool(bold))}"/><w:i w:val="0"/><w:iCs w:val="0"/>'
    rpr += f'<w:caps w:val="{int(bool(caps))}"/>'
    rpr += f'<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/><w:lang w:val="{lang}" w:eastAsia="{lang}" w:bidi="ar-SA"/>'
    head = f'<w:name w:val="{name}"/>'
    head += f'<w:next w:val="{next_id}"/>' if next_id else ""
    head += f'<w:uiPriority w:val="{ui}"/>' if ui is not None else ""
    head += "<w:qFormat/>"
    return parse_xml(f'<w:style {nsdecls("w")} w:type="paragraph"{" w:customStyle=\"1\"" if custom else ""} w:styleId="{sid}">'
                     f"{head}<w:pPr>{ppr}</w:pPr><w:rPr>{rpr}</w:rPr></w:style>")


def _builtin_id(styles_el, name, fallback):
    """id встроенного стиля основы по его имени («heading 1», «toc 2»): в русских шаблонах id бывает «1», «11»…"""
    for st in styles_el.findall(qn("w:style")):
        nm = st.find(qn("w:name"))
        if st.get(qn("w:type")) == "paragraph" and nm is not None and (nm.get(qn("w:val")) or "").lower() == name:
            return st.get(qn("w:styleId"))
    return fallback


def install_styles(doc, style, text_w):
    """Поставить в основу стили RK* и переопределить Heading 1–3 и TOC 1–3. Возвращает id стилей заголовков и оглавления."""
    styles_el = doc.styles.element
    body, hs = style["body"], style["headings"]
    font, size = body["font"], body["size_pt"]
    lang = {"ru": "ru-RU", "en": "en-US"}.get(style.get("language", "ru"), "ru-RU")
    line = int(round(240 * float(body.get("line_spacing", 1.5))))
    indent = _tw_cm(body.get("first_indent_cm", 1.25))
    lst, tab, fig = style["listing"], style["table"], style["figure"]
    ids = {}
    new = [
        _style_xml(STYLE_BODY, "RK Body", font, size, lang, align=body.get("align", "justify"), first=indent, line=line),
        _style_xml(STYLE_LIST, "RK List", font, size, lang, align=body.get("align", "justify"), first=indent, line=line),
        _style_xml(STYLE_CAPTION_FIG, "RK Caption Figure", font, fig.get("size_pt", size), lang, align=fig.get("align", "center"),
                   line=line, after=240, keep_lines=True),
        _style_xml(STYLE_CAPTION_TAB, "RK Caption Table", font, size, lang, align=tab.get("caption_align", "left"), line=line,
                   before=120, keep_next=True, keep_lines=True),
        _style_xml(STYLE_CAPTION_LST, "RK Caption Listing", font, size, lang, align="left", line=line, before=120, keep_next=True,
                   keep_lines=True),
        _style_xml(STYLE_LISTING, "RK Listing", lst.get("font", "Courier New"), lst.get("size_pt", 10), lang, align="left",
                   line=int(round(240 * float(lst.get("line_spacing", 1.0))))),
        _style_xml(STYLE_TABLE_TEXT, "RK Table Text", font, tab.get("size_pt", 12), lang, align="left"),
        _style_xml(STYLE_FORMULA, "RK Formula", font, size, lang, align="center", before=120, after=120),
        _style_xml(STYLE_TITLE, "RK Title", font, size, lang, align="center"),
    ]
    for n in (1, 2, 3):
        h = hs.get(f"h{n}") or {}
        ids[f"h{n}"] = sid = _builtin_id(styles_el, f"heading {n}", f"Heading{n}")
        new.append(_style_xml(sid, f"heading {n}", font, h.get("size_pt", size), lang, align=h.get("align", "left"),
                              first=indent if h.get("indent", True) and h.get("align", "left") != "center" else 0, line=line,
                              before=240, after=240 if n == 1 else 120, bold=h.get("bold", True), caps=h.get("caps", False),
                              keep_next=hs.get("keep_with_next", True), keep_lines=True, page_break=h.get("new_page", False),
                              outline=n - 1, next_id=STYLE_BODY, ui=9, custom=False))
    for n in (1, 2, 3):
        ids[f"toc{n}"] = sid = _builtin_id(styles_el, f"toc {n}", f"TOC{n}")
        new.append(_style_xml(sid, f"toc {n}", font, size, lang, align="left", left=(n - 1) * _tw_cm(0.5), line=line,
                              next_id=STYLE_BODY, ui=39, tab_right=text_w, custom=False))
    for el in new:
        sid = el.get(qn("w:styleId"))
        old = next((s for s in styles_el.findall(qn("w:style")) if s.get(qn("w:styleId")) == sid), None)
        if old is not None:
            styles_el.replace(old, el)
        else:
            styles_el.append(el)
    return ids


# ---------- чтение файлов для таблиц и листингов ----------

def _ranges(spec, what, line):
    """«1,3,5-7» → [1, 3, 5, 6, 7] (номера с 1)."""
    if isinstance(spec, (list, tuple)):
        return [int(x) for x in spec]
    out = []
    for part in str(spec).replace(" ", "").split(","):
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not m or int(m.group(1)) < 1 or (m.group(2) and int(m.group(2)) < int(m.group(1))):
            raise KitError(f"{what}: не разобрать «{spec}»", line=line, hint="пиши номера с 1 через запятую и диапазоны через дефис, например 1,3,5-7")
        out.extend(range(int(m.group(1)), int(m.group(2) or m.group(1)) + 1))
    return out


def _read_text(path, line):
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return clean_text(raw.decode(enc))
        except UnicodeDecodeError:
            continue
    raise KitError(f"файл не читается как текст: {os.path.basename(path)}", line=line, hint="сохрани файл в UTF-8")


def load_csv_table(path, cols=None, rows=None, header=None, line=None):
    """Таблица из файла: первая строка файла — шапка, разделитель «;» или табуляция. Возвращает (шапка, строки)."""
    name = os.path.basename(path)
    try:
        with open(path, "rb") as f:
            text = clean_text(f.read().decode("utf-8-sig"))
    except UnicodeDecodeError:
        raise KitError(f"файл таблицы {name} не в UTF-8", line=line, hint="сохраняй таблицу так: open(путь, 'w', encoding='utf-8', newline='')")
    lines = [ln for ln in text.split("\n") if ln.strip()]
    if not lines:
        raise KitError(f"файл таблицы {name} пуст", line=line)
    delim = ";" if ";" in lines[0] else "\t" if "\t" in lines[0] else None
    if delim is None and "," in lines[0]:
        raise KitError(f"в файле таблицы {name} разделитель — запятая", line=line,
                       hint="запиши файл с разделителем «;»: csv.writer(f, delimiter=';') — запятая нужна для десятичных дробей")
    data = [[c.strip() for c in r] for r in csv.reader(lines, delimiter=delim or ";")]
    width = len(data[0])
    for k, r in enumerate(data, 1):
        if len(r) != width:
            raise KitError(f"в файле таблицы {name} в строке {k} ячеек {len(r)}, а в шапке {width}", line=line, hint="выровняй число ячеек в строках файла")
    head, body = data[0], data[1:]
    if rows:
        want = _ranges(rows, "rows", line)
        body = [body[i - 1] for i in want if i <= len(body)]
    if cols:
        want = _ranges(cols, "cols", line)
        bad = [c for c in want if c > width]
        if bad:
            raise KitError(f"cols: столбца {bad[0]} нет — в файле {name} столбцов {width}", line=line)
        head, body = [head[c - 1] for c in want], [[r[c - 1] for c in want] for r in body]
    if header:
        names = [h.strip() for h in header.split(";")] if isinstance(header, str) else list(header)
        if len(names) != len(head):
            raise KitError(f"header: названий {len(names)}, а столбцов в таблице {len(head)}", line=line, hint="в header перечисли названия всех выбранных столбцов через «;»")
        head = names
    return head, body


def load_listing_file(path, from_=None, to=None, lines=None, line=None):
    """Строки листинга из файла: целиком, по номерам строк (lines=10-40) или от строки с from по строку с to включительно."""
    src = _read_text(path, line).split("\n")
    if src and src[-1] == "":
        src.pop()
    name = os.path.basename(path)
    if lines:
        if isinstance(lines, (list, tuple)) and len(lines) == 2:
            lines = f"{lines[0]}-{lines[1]}"  # mdparse отдаёт диапазон парой [с, по]
        m = re.fullmatch(r"\s*(\d+)\s*(?:-\s*(\d+))?\s*", str(lines))
        if not m or int(m.group(1)) < 1 or int(m.group(2) or m.group(1)) < int(m.group(1)):
            raise KitError(f"lines: не разобрать «{lines}»", line=line, hint="пиши lines=10-40 (номера строк файла с 1)")
        out = src[int(m.group(1)) - 1:int(m.group(2) or m.group(1))]
        if not out:
            raise KitError(f"lines={lines}: в файле {name} только {len(src)} строк", line=line)
        return out
    start, end = 0, len(src) - 1
    if from_:
        start = next((i for i, s in enumerate(src) if from_ in s), None)
        if start is None:
            raise KitError(f"from: в файле {name} нет строки с «{from_}»", line=line, hint="скопируй подстроку из файла буква в букву или укажи lines=N-M")
    if to:
        end = next((i for i in range(start + 1, len(src)) if to in src[i]), None)
        if end is None:
            raise KitError(f"to: в файле {name} после начала нет строки с «{to}»", line=line, hint="скопируй подстроку из файла буква в букву или укажи lines=N-M")
    return src[start:end + 1]


# ---------- сборка ----------

def _plain(runs):
    """Текст строчных элементов без оформления (для длины ячеек и сравнения названий)."""
    if isinstance(runs, str):
        return runs
    return "".join(r if isinstance(r, str) else str(r.get("t", r.get("math", r.get("ref", "")))) for r in runs or [])


def _each_runs(b):
    """Все списки строчных элементов блока (чтобы найти в них ссылки)."""
    t = b.get("type")
    if t == "heading":
        yield b.get("text_runs")
    elif t == "para":
        yield b.get("runs")
    elif t == "list":
        for it in b.get("items") or []:
            yield it.get("runs")
    elif t == "figure":
        yield b.get("caption_runs")
    elif t == "table":
        yield b.get("caption_runs")
        yield from b.get("header") or []
        if isinstance(b.get("rows"), list) and not b.get("file"):
            for row in b["rows"]:
                yield from row
    elif t == "listing":
        yield b.get("caption")


def _soft_wrap(text):
    """Мягкий перенос после «_», «/», «.» в длинных словах без пробелов (пути, имена) — чтобы ячейка не распирала таблицу."""
    def fix(m):
        w = m.group(0)
        return w if re.fullmatch(r"[-+\d.,eE%]+", w) else re.sub(r"(?<=[_/.])(?=[^\W_])", SOFT, w)
    return re.sub(r"\S{13,}", fix, text)


class _Builder:
    """Состояние одной сборки: документ, профиль, счётчики и номера объектов."""

    def __init__(self, doc, style, base, ids, metrics):
        self.doc, self.style, self.base, self.ids = doc, style, base, ids
        self.page_w, self.page_h, self.mar, self.text_w, self.text_h = metrics
        self.warnings = []
        self.counts = {"sections": 0, "figures": 0, "tables": 0, "formulas": 0, "listings": 0}
        self.labels, self.numbers = {}, {}  # метка → номер; индекс блока → номер
        self.h = [0, 0, 0]
        self.code_font = style["listing"].get("font", "Courier New")
        self.comma = style.get("decimal", "comma") == "comma"
        self.gap = False  # предыдущим блоком была таблица или листинг: следующему абзацу — отбивка сверху
        self.no_break = False  # первый блок тела без титульного листа не уходит на новую страницу
        self.in_sources = False  # идёт раздел списка источников (заголовок 1-го уровня из sources.title или «СПИСОК …»)

    # --- номера и ссылки ---
    def number(self, blocks):
        """Первый проход: номера рисунков, таблиц, листингов и формул; метки; проверка ссылок.
        Предупреждения (метка без ссылки, таблица без подписи) даёт mdparse — здесь только ошибки."""
        n = {"fig": 0, "tbl": 0, "lst": 0, "eq": 0}
        prefix = {"figure": "fig", "table": "tbl", "listing": "lst", "formula": "eq"}
        lines = {}
        for i, b in enumerate(blocks):
            kind = prefix.get(b.get("type"))
            if not kind:
                continue
            numbered = (kind == "fig" or (kind == "tbl" and bool(b.get("caption_runs"))) or (kind == "lst" and bool(b.get("caption")))
                        or (kind == "eq" and b.get("numbered", True) and self.style["formula"].get("numbered", True)))
            if numbered:
                n[kind] += 1
                self.numbers[i] = n[kind]
            label = b.get("id")
            if not label:
                continue
            label = label if ":" in label else f"{kind}:{label}"
            if label in self.labels or label in lines:
                raise KitError(f"метка «{label}» повторяется (первая — в строке {lines.get(label)})", line=b.get("line"), hint="дай объектам разные метки")
            lines[label] = b.get("line")
            if numbered:
                self.labels[label] = f"({n[kind]})" if kind == "eq" else str(n[kind])
        for b in blocks:
            for runs in _each_runs(b):
                for r in runs if isinstance(runs, list) else []:
                    if isinstance(r, dict) and "ref" in r:
                        if r["ref"] not in self.labels:
                            why = "у этого объекта нет номера (нет подписи или стоит {-})" if r["ref"] in lines else "объекта с такой меткой нет"
                            raise KitError(f"ссылка @{r['ref']}: {why}", line=b.get("line"),
                                           hint="метка ставится в хвосте объекта, например {#fig:имя}; проверь написание" + (f"; есть метки: {', '.join(sorted(lines)[:8])}" if lines else ""))

    # --- абзацы и строчные элементы ---
    def p(self, sid, container=None):
        p = (container or self.doc).add_paragraph()
        p._p.get_or_add_pPr().style = sid
        if self.gap and container is None and sid in (STYLE_BODY, STYLE_LIST):
            p.paragraph_format.space_before = Pt(6)
        if container is None:
            self.gap = False
        return p

    def math(self, latex, display, line):
        try:
            from .mathconv import latex_to_omml
        except ImportError as e:
            raise KitError(f"модуль формул не загрузился: {e}", line=line, code=EXIT_SELF)
        try:
            el = copy.deepcopy(latex_to_omml(clean_text(latex).strip(), display))
        except MathError as e:
            at = f" (позиция {e.pos})" if e.pos is not None else ""
            raise KitError(f"формула не разобрана: {e.message}{at}: {latex}", line=line, hint=e.hint)
        is_para = el.tag == f"{{{_M}}}oMathPara"
        if display and not is_para:
            wrap = etree.Element(f"{{{_M}}}oMathPara", nsmap={"m": _M})
            wrap.append(el)
            return wrap
        if not display and is_para:
            return el.find(f"{{{_M}}}oMath")
        return el

    def runs(self, p, runs, line, upper=False, soft=False, bold=False):
        """Строчные элементы → прогоны абзаца: текст (b, i, code), формула в строке, номер по ссылке."""
        for r in [runs] if isinstance(runs, str) else runs or []:
            r = {"t": r} if isinstance(r, str) else r
            if "math" in r:
                p._p.append(self.math(r["math"], False, line))
                continue
            text = self.labels[r["ref"]] if "ref" in r else clean_text(r.get("t", "")).replace("\n", " ")
            if upper and not r.get("code"):
                text = text.upper()
            if soft:
                text = _soft_wrap(text)
            if not text:
                continue
            parts = text.split(SOFT)
            run = p.add_run(parts[0])
            for part in parts[1:]:  # знак U+00AD Word печатает как дефис всегда; его «мягкий перенос» — элемент w:softHyphen
                run._r.append(parse_xml(f"<w:softHyphen {nsdecls('w')}/>"))
                t = parse_xml(f"<w:t {nsdecls('w')}/>")
                t.text = part
                run._r.append(t)
            if r.get("b") or bold:
                run.bold = True
            if r.get("i"):
                run.italic = True
            if r.get("code"):
                fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
                for a in ("ascii", "hAnsi", "cs", "eastAsia"):
                    fonts.set(qn("w:" + a), self.code_font)

    # --- блоки ---
    def heading(self, b):
        level = int(b.get("level", 1))
        if level not in (1, 2, 3):
            raise KitError(f"заголовок уровня {level} не поддерживается", line=b.get("line"), hint="используй #, ## или ###")
        hs = self.style["headings"]
        h = hs.get(f"h{level}") or {}
        structural = bool(b.get("structural")) and level == 1
        toc_title = b.get("_toc_title")
        if level == 1:
            name = _plain(b.get("text_runs")).strip().upper()
            self.in_sources = name == str((self.style.get("sources") or {}).get("title") or "").upper() or name.startswith("СПИСОК ")
        if level == 1 and not toc_title:
            self.counts["sections"] += 1
        p = self.p(STYLE_BODY if toc_title else self.ids[f"h{level}"])
        fmt = p.paragraph_format
        prefix, caps, bold = "", h.get("caps", False), False
        if structural or toc_title:
            st = hs.get("structural") or {}
            caps = st.get("caps", True)
            fmt.alignment = ALIGN.get(st.get("align", "center"), WD_ALIGN_PARAGRAPH.CENTER)
            if st.get("align", "center") == "center":
                fmt.first_line_indent = Twips(0)
            if bool(st.get("new_page", True)) != bool(h.get("new_page", False)) or toc_title:
                fmt.page_break_before = bool(st.get("new_page", True))
            bold = st.get("bold", True) and (toc_title or not h.get("bold", True))
            if toc_title:
                fmt.keep_with_next = True
                fmt.space_before, fmt.space_after = Pt(12), Pt(12)
        elif b.get("numbered", True) and hs.get("numbering", "1 Название") != "none":
            self.h[level - 1] += 1
            for k in range(level, 3):
                self.h[k] = 0
            num = ".".join(str(max(x, 1)) for x in self.h[:level])
            mode = hs.get("numbering")
            if mode not in ("1 Название", "1. Название"):
                raise KitError(f"style.json: headings.numbering = {mode!r} не поддерживается", hint='допустимо: "1 Название", "1. Название", "none"')
            prefix = num + (". " if mode == "1. Название" else " ")
        if self.no_break:
            fmt.page_break_before = False
        if prefix:
            p.add_run(prefix)
        self.runs(p, b.get("text_runs"), b.get("line"), upper=caps, bold=bold)

    def para(self, b, keep_next=False):
        p = self.p(STYLE_BODY)
        self.runs(p, b.get("runs"), b.get("line"))
        # абзац, который вводит рисунок («…показано на рисунке 5»), не остаётся один внизу страницы, когда рисунок ушёл на следующую
        if keep_next:
            p.paragraph_format.keep_with_next = True

    def list_block(self, b):
        ls = self.style["list"]
        step = _tw_cm(self.style["body"].get("first_indent_cm", 1.25))
        for k, it in enumerate(b.get("items") or []):
            ordered = it.get("ordered", b.get("ordered"))
            p = self.p(STYLE_BODY if ordered else STYLE_LIST)
            if it.get("level"):
                p.paragraph_format.left_indent = Twips(step * int(it["level"]))
            if ordered:
                n = it.get("number") if it.get("number") is not None else k + 1
                # список источников нумеруется «1.» при любом list.number_format: в тексте на него ссылаются как [1]
                mark = f"{n}." if self.in_sources or str(ls.get("number_format", "1)")).endswith(".") else f"{n})"
            else:
                mark = ls.get("bullet", "–")
            p.add_run(mark + " ")  # неразрывный пробел: маркер не отрывается и не растягивается при выравнивании по ширине
            self.runs(p, it.get("runs"), it.get("line") or b.get("line"))

    def figure(self, b, number):
        from PIL import Image
        line = b.get("line")
        paths = [resolve_path(self.base, x, line) for x in b.get("images") or []]
        if not paths:
            raise KitError("у рисунка нет файла картинки", line=line)
        sizes = []
        for path in paths:
            if os.path.splitext(path)[1].lower() not in IMAGE_EXT:
                raise KitError(f"формат рисунка не поддерживается: {os.path.basename(path)}", line=line, hint="сохрани картинку в png или jpg")
            try:
                with Image.open(path) as im:
                    sizes.append(im.size)
            except Exception:
                raise KitError(f"рисунок не открывается: {os.path.basename(path)}", line=line, hint="файл повреждён или это не картинка — пересоздай его")
        n = len(paths)
        widths = list(b.get("widths") or [])
        if len(widths) == 1:
            widths = widths * n
        if widths and len(widths) != n:
            raise KitError(f"width: значений {len(widths)}, а картинок {n}", line=line, hint="укажи одну ширину на все или по одной на каждую: width=7cm,7cm")
        row = bool(b.get("row")) and n > 1
        max_w, max_h = self.text_w * EMU_TWIP, self.text_h * EMU_TWIP - 2 * EMU_CM
        room = (max_w - (n - 1) * ROW_GAP) // n if row else max_w
        dims = []
        for k, (pw, ph) in enumerate(sizes):
            w = int(float(widths[k]) * EMU_CM) if widths else room
            if w > max_w:
                self.warnings.append(f"рисунок {number}: ширина {widths[k]} см больше ширины текста — уменьшена (строка {line})")
                w = max_w
            hgt = int(w * ph / max(pw, 1))
            if hgt > max_h:
                w, hgt = int(w * max_h / hgt), max_h
            dims.append([w, hgt])
        if row:
            total = sum(d[0] for d in dims) + (n - 1) * ROW_GAP
            if total > max_w:
                k = (max_w - (n - 1) * ROW_GAP) / (total - (n - 1) * ROW_GAP)
                dims = [[int(w * k), int(hgt * k)] for w, hgt in dims]
        p = None
        for k, path in enumerate(paths):
            if p is None or not row:
                p = self.p(STYLE_CAPTION_FIG)  # абзац рисунка: тот же стиль, что у подписи, но интервал 1,0 и без отбивки
                fmt = p.paragraph_format
                fmt.keep_with_next, fmt.line_spacing, fmt.space_before, fmt.space_after = True, 1.0, Pt(6), Pt(0)
                fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif row:
                p.add_run(" ")
            p.add_run().add_picture(path, width=Emu(dims[k][0]), height=Emu(dims[k][1]))
        cap = self.p(STYLE_CAPTION_FIG)
        fig = self.style["figure"]
        cap.add_run(f"{fig.get('word', 'Рисунок')} {number}" + (fig.get("sep", " – ") if b.get("caption_runs") else ""))
        self.runs(cap, b.get("caption_runs"), line)
        self.counts["figures"] += 1

    def _table(self, n_rows, widths, borders, mark=None):
        """Таблица заданной ширины столбцов (твипы): по центру, фиксированная раскладка, строки не разрываются."""
        tbl = self.doc.add_table(rows=n_rows, cols=len(widths))
        total = sum(widths)
        side = ('w:val="single" w:sz="4" w:space="0" w:color="000000"' if borders else 'w:val="nil"')
        xml = (f'<w:tblPr {nsdecls("w")}><w:tblW w:w="{total}" w:type="dxa"/><w:jc w:val="center"/><w:tblBorders>'
               + "".join(f"<w:{s} {side}/>" for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
               + '</w:tblBorders><w:tblLayout w:type="fixed"/><w:tblCellMar><w:left w:w="85" w:type="dxa"/><w:right w:w="85" w:type="dxa"/>'
               + '</w:tblCellMar><w:tblLook w:val="0000"/>' + (f'<w:tblCaption w:val="{mark}"/>' if mark else "") + "</w:tblPr>")
        tbl._tbl.replace(tbl._tbl.tblPr, parse_xml(xml))
        for k, col in enumerate(tbl._tbl.tblGrid.findall(qn("w:gridCol"))):
            col.set(qn("w:w"), str(widths[k]))
        for tr in tbl._tbl.findall(qn("w:tr")):
            tr.get_or_add_trPr().append(parse_xml(f"<w:cantSplit {nsdecls('w')}/>"))
            for k, tc in enumerate(tr.findall(qn("w:tc"))):
                tcw = tc.get_or_add_tcPr().get_or_add_tcW()
                tcw.set(qn("w:w"), str(widths[k]))
                tcw.set(qn("w:type"), "dxa")
        return tbl

    def _cell(self, cell, content, line, align, bold=False):
        p = cell.paragraphs[0]
        p._p.get_or_add_pPr().style = STYLE_TABLE_TEXT
        if align and align not in ("l", "left"):
            p.paragraph_format.alignment = ALIGN.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        text = _plain(content).strip()
        if self.comma and _DECIMAL.match(text) and all(isinstance(r, str) or "t" in r for r in ([content] if isinstance(content, str) else content)):
            content = [{"t": text.replace(".", ",")}]
        self.runs(p, content, line, soft=True, bold=bold)

    def table(self, b, number):
        line = b.get("line")
        if b.get("file"):
            want = b.get("rows")
            if isinstance(want, (list, tuple)) and len(want) == 2:
                want = f"{want[0]}-{want[1]}"  # mdparse отдаёт диапазон парой [с, по]
            header, rows = load_csv_table(resolve_path(self.base, b["file"], line), b.get("cols"), want, b.get("header"), line)
        else:
            header, rows = b.get("header") or [], b.get("rows") or []
        n_cols = len(header) if header else (len(rows[0]) if rows else 0)
        if not n_cols:
            raise KitError("таблица пуста", line=line, hint="в таблице должна быть шапка и хотя бы одна строка")
        for k, row in enumerate(rows, 1):
            if len(row) != n_cols:
                raise KitError(f"в строке {k} таблицы ячеек {len(row)}, а в шапке {n_cols}", line=line, hint="число ячеек в каждой строке должно совпадать с шапкой")
        shares = b.get("widths")
        if shares and len(shares) != n_cols:
            raise KitError(f"widths: долей {len(shares)}, а столбцов {n_cols}", line=line, hint="перечисли доли всех столбцов: widths=3,2,2")
        if not shares:  # ширины по самому длинному содержимому столбца, в разумных пределах
            # +3 знака — поля ячейки: иначе в таблице на 9 столбцов короткое слово («норма») рвётся посреди
            shares = [min(max([len(_plain(r[c])) for r in ([header] if header else []) + rows] + [3]), 40) + 3 for c in range(n_cols)]
        widths = [int(self.text_w * float(s) / sum(float(x) for x in shares)) for s in shares]
        widths[-1] = self.text_w - sum(widths[:-1])
        tab = self.style["table"]
        if b.get("caption_runs"):
            cap = self.p(STYLE_CAPTION_TAB)
            cap.add_run(f"{tab.get('word', 'Таблица')} {number}{tab.get('sep', ' – ')}")
            self.runs(cap, b["caption_runs"], line)
        tbl = self._table(len(rows) + (1 if header else 0), widths, borders=True)
        align = list(b.get("align") or []) + ["l"] * n_cols
        trs = tbl.rows
        if header:
            trs[0]._tr.get_or_add_trPr().append(parse_xml(f"<w:tblHeader {nsdecls('w')}/>"))
            for c, cell in enumerate(trs[0].cells):
                self._cell(cell, header[c], line, "c", bold=tab.get("header_bold", False))
        for k, row in enumerate(rows):
            for c, cell in enumerate(trs[k + (1 if header else 0)].cells):
                self._cell(cell, row[c], line, align[c])
        self.counts["tables"] += 1
        self.gap = True

    def listing(self, b, number):
        line = b.get("line")
        if b.get("file"):
            lines = load_listing_file(resolve_path(self.base, b["file"], line), b.get("from"), b.get("to"), b.get("lines"), line)
        else:
            lines = b.get("lines") or []
            lines = lines.split("\n") if isinstance(lines, str) else lines
        if b.get("caption"):
            lst = self.style["listing"]
            cap = self.p(STYLE_CAPTION_LST)
            cap.add_run(f"{lst.get('word', 'Листинг')} {number}{lst.get('sep', ' – ')}")
            self.runs(cap, b["caption"], line)
        p = None
        for k, text in enumerate(lines or [""]):
            p = self.p(STYLE_LISTING)
            if k == 0 and not b.get("caption"):
                p.paragraph_format.space_before = Pt(6)
            text = clean_text(text).replace("\n", " ").replace("\t", "    ")
            if text:
                run = p.add_run(text)
                run._r.find(qn("w:t")).set(_XML_SPACE, "preserve")  # пробелы кода — как в исходнике
        p.paragraph_format.space_after = Pt(6)
        self.counts["listings"] += 1
        self.gap = True

    def formula(self, b, number):
        line = b.get("line")
        omml = self.math(b.get("latex", ""), True, line)
        side = int(self.text_w * 0.1)
        tbl = self._table(1, [side, self.text_w - 2 * side, side], borders=False, mark=EQ_TABLE_MARK)
        cells = tbl.rows[0].cells
        for k, cell in enumerate(cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p._p.get_or_add_pPr().style = STYLE_FORMULA
            if k == 1:
                p._p.append(omml)
            elif k == 2:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if number:
                    p.add_run(f"({number})")
        self.counts["formulas"] += 1
        self.gap = False

    def toc(self, b, with_title):
        toc = self.style.get("toc") or {}
        if with_title and toc.get("title"):
            self.heading({"type": "heading", "level": 1, "line": b.get("line"), "text_runs": [{"t": toc["title"]}], "_toc_title": True})
        p = self.p(self.ids["toc1"])
        levels = min(max(int(toc.get("levels", 2)), 1), 3)
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'))
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> TOC \\o "1-{levels}" \\h \\z \\u </w:instrText></w:r>'))
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>'))
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:t>{TOC_STUB}</w:t></w:r>'))
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'))

    def pagebreak(self, b):
        p = self.p(STYLE_BODY)
        p.paragraph_format.first_line_indent = Twips(0)
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:br w:type="page"/></w:r>'))

    def write(self, blocks):
        blocks = [dict(b) for b in blocks]
        self.number(blocks)
        # заголовок «СОДЕРЖАНИЕ» прямо перед [TOC] — название оглавления, а не раздел: в само оглавление он попадать не должен
        toc_names = {"СОДЕРЖАНИЕ", "ОГЛАВЛЕНИЕ", str((self.style.get("toc") or {}).get("title") or "").upper()}
        own_title = set()
        for i, b in enumerate(blocks):
            if b.get("type") == "toc" and i and blocks[i - 1].get("type") == "heading" and _plain(blocks[i - 1].get("text_runs")).strip().upper() in toc_names:
                blocks[i - 1]["_toc_title"] = True
                own_title.add(i)
        for i, b in enumerate(blocks):
            t = b.get("type")
            if t == "heading":
                self.heading(b)
            elif t == "para":
                self.para(b, keep_next=i + 1 < len(blocks) and blocks[i + 1].get("type") == "figure")
            elif t == "list":
                self.list_block(b)
            elif t == "figure":
                self.figure(b, self.numbers[i])
            elif t == "table":
                self.table(b, self.numbers.get(i))
            elif t == "listing":
                self.listing(b, self.numbers.get(i))
            elif t == "formula":
                self.formula(b, self.numbers.get(i))
            elif t == "toc":
                self.toc(b, with_title=i not in own_title)
            elif t == "pagebreak":
                self.pagebreak(b)
            else:
                raise KitError(f"неизвестный вид блока: {t!r}", line=b.get("line"), code=EXIT_SELF)
            self.no_break = False


# ---------- лист, номер страницы ----------

def _setup_page(doc, metrics):
    """Лист и поля — явно в разделе тела (последний sectPr документа; разделы титула шаблона не трогаются)."""
    w, h, mar, _, _ = metrics
    doc.element.body.get_or_add_sectPr()
    sec = doc.sections[-1]
    sec.page_width, sec.page_height = Twips(w), Twips(h)
    pg = sec._sectPr.find(qn("w:pgSz"))
    if w > h:
        pg.set(qn("w:orient"), "landscape")
    elif pg.get(qn("w:orient")) is not None:
        del pg.attrib[qn("w:orient")]
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = (Twips(mar[k]) for k in ("left", "right", "top", "bottom"))
    sec.header_distance = sec.footer_distance = Twips(709)
    sec.gutter = Twips(0)
    return sec


def _has_page_field(doc):
    """В колонтитулах основы уже стоит поле PAGE?"""
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/header") or rel.reltype.endswith("/footer"):
            el = rel.target_part.element
            instr = " ".join((x.text or "") for x in el.iter(qn("w:instrText"))) + " " + " ".join(x.get(qn("w:instr"), "") for x in el.iter(qn("w:fldSimple")))
            if re.search(r"\bPAGE\b", instr):
                return True
    return False


def _page_number(doc, style, first_page_special):
    """Поле PAGE в колонтитуле раздела тела. first_page_special — на первой странице (титуле) номер не показывать."""
    pn = style.get("page_number") or {}
    pos = pn.get("position", "bottom-center")
    if pos == "none" or _has_page_field(doc):
        return
    if pos not in ("bottom-center", "bottom-right", "top-right"):
        raise KitError(f"style.json: page_number.position = {pos!r} не поддерживается", hint="допустимо: bottom-center, bottom-right, top-right, none")
    sec = doc.sections[-1]
    part = sec.footer if pos.startswith("bottom") else sec.header
    part.is_linked_to_previous = False
    p = part.paragraphs[0] if part.paragraphs and not part.paragraphs[0].text.strip() else part.add_paragraph()
    font, sz = style["body"]["font"], int(round(float(pn.get("size_pt", 12)) * 2))
    rpr = f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}" w:eastAsia="{font}"/><w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/></w:rPr>'
    ppr = (f'<w:pPr {nsdecls("w")}><w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/><w:ind w:left="0" w:right="0" w:firstLine="0"/>'
           f'<w:jc w:val="{"center" if pos.endswith("center") else "right"}"/></w:pPr>')
    for old in list(p._p):
        p._p.remove(old)  # в заготовке колонтитула python-docx стоит стиль Footer/Header, которого может не быть в основе
    p._p.append(parse_xml(ppr))
    for inner in ('<w:fldChar w:fldCharType="begin"/>', '<w:instrText xml:space="preserve"> PAGE </w:instrText>',
                  '<w:fldChar w:fldCharType="separate"/>', "<w:t>1</w:t>", '<w:fldChar w:fldCharType="end"/>'):
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}>{rpr}{inner}</w:r>'))
    if first_page_special:
        sec.different_first_page_header_footer = True


# ---------- завершение и самопроверка ----------

def finalize(doc, out_path):
    """Сохранить документ так, чтобы тот же вход давал побайтно тот же файл; затем самопроверка (код 5)."""
    cp = doc.core_properties
    cp.created = cp.modified = FIXED_DATE
    cp.last_modified_by, cp.revision = "", 1
    buf = io.BytesIO()
    doc.save(buf)
    src = zipfile.ZipFile(buf)
    data = {n: src.read(n) for n in src.namelist() if not n.lower().startswith("docprops/thumbnail")}
    if len(data) != len(src.namelist()) and "_rels/.rels" in data:  # миниатюру шаблона убираем вместе со связью на неё
        rels = etree.fromstring(data["_rels/.rels"])
        for rel in list(rels):
            if "thumbnail" in (rel.get("Target") or "").lower():
                rels.remove(rel)
        data["_rels/.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
    first = [n for n in ("[Content_Types].xml", "_rels/.rels") if n in data]
    tmp = out_path + ".tmp"
    try:
        os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
            for name in first + sorted(n for n in data if n not in first):
                info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
                info.compress_type, info.create_system, info.external_attr = zipfile.ZIP_DEFLATED, 0, 0
                z.writestr(info, data[name])
        problems = self_check(tmp)
        if problems:
            raise KitError("собранный файл не прошёл самопроверку: " + "; ".join(problems[:5]), code=EXIT_SELF,
                           hint="это ошибка сборщика, а не report.md: сообщи о ней и приложи report.md")
        os.replace(tmp, out_path)
    except PermissionError:
        raise KitError(f"не удаётся записать {out_path}: файл занят", code=EXIT_WORD, hint="закрой отчёт в Word (или укажи другое имя через --out) и повтори сборку")
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


def self_check(path):
    """Проверка собранного DOCX: открывается, связи целы, стили существуют, id рисунков уникальны, разметка не протекла.
    Возвращает список найденных проблем (пустой — всё в порядке)."""
    problems = []
    try:
        z = zipfile.ZipFile(path)
        names = set(z.namelist())
    except Exception as e:
        return [f"файл не открывается: {type(e).__name__}: {e}"]
    try:
        Document(path)
    except Exception as e:
        problems.append(f"файл не открывается в python-docx: {type(e).__name__}: {e}")
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    styles = etree.fromstring(z.read("word/styles.xml")) if "word/styles.xml" in names else etree.Element("x")
    style_ids = {s.get(W + "styleId") for s in styles.iter(W + "style")}
    code_font = None
    for s in styles.iter(W + "style"):
        if s.get(W + "styleId") == STYLE_LISTING:
            rf = s.find(f"{W}rPr/{W}rFonts")
            code_font = rf.get(W + "ascii") if rf is not None else None
    for name in sorted(names):
        m = re.fullmatch(r"word/((?:document|header\d*|footer\d*|footnotes|endnotes)\.xml)", name)
        if not m:
            continue
        root = etree.fromstring(z.read(name))
        rels = {}
        rels_name = f"word/_rels/{m.group(1)}.rels"
        if rels_name in names:
            for rel in etree.fromstring(z.read(rels_name)):
                rels[rel.get("Id")] = (rel.get("Target") or "", rel.get("TargetMode"))
        for el in root.iter():
            for attr in ("embed", "id", "link", "pict"):
                rid = el.get(f"{{{_R}}}{attr}")
                if rid is None:
                    continue
                if rid not in rels:
                    problems.append(f"{name}: связь {rid} не описана в {rels_name}")
                elif rels[rid][1] != "External":
                    target = rels[rid][0]
                    full = target.lstrip("/") if target.startswith("/") else posixpath.normpath("word/" + target)
                    if full not in names:
                        problems.append(f"{name}: связь {rid} ведёт на отсутствующую часть {full}")
            if el.tag in (W + "pStyle", W + "rStyle", W + "tblStyle") and el.get(W + "val") not in style_ids:
                problems.append(f"{name}: стиль «{el.get(W + 'val')}» не описан в styles.xml")
        if m.group(1) != "document.xml":
            continue
        ids = [el.get("id") for el in root.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")]
        dup = sorted({i for i in ids if ids.count(i) > 1})
        if dup:
            problems.append(f"повторяются id рисунков (wp:docPr): {', '.join(dup[:5])}")
        for p in root.iter(W + "p"):
            ps = p.find(f"{W}pPr/{W}pStyle")
            if ps is not None and ps.get(W + "val") == STYLE_LISTING:
                continue  # в листинге «$$» или «![» — обычный код
            text = ""
            for r in p.iter(W + "r"):
                rf = r.find(f"{W}rPr/{W}rFonts")
                if code_font and rf is not None and rf.get(W + "ascii") == code_font:
                    continue  # `код` в строке
                text += "".join(t.text or "" for t in r.findall(W + "t"))
            for mark in LEFTOVERS:
                if mark in text:
                    problems.append(f"в тексте осталась разметка «{mark}»: {text[:60]}")
    return problems


def _run_post(doc, base, post):
    """--post: функция post(doc) из файла пользователя над готовым документом (до сохранения)."""
    path = resolve_path(base, post)
    ns = {"__name__": "rk_post", "__file__": path}
    try:
        with open(path, encoding="utf-8-sig") as f:
            exec(compile(f.read(), path, "exec"), ns)
        if not callable(ns.get("post")):
            raise KitError(f"в {os.path.basename(path)} нет функции post(doc)", hint="объяви def post(doc): … — она получает готовый python-docx Document")
        ns["post"](doc)
    except KitError:
        raise
    except Exception as e:
        raise KitError(f"{os.path.basename(path)} завершился ошибкой: {type(e).__name__}: {e}", hint="исправь post.py или убери --post")


def build_docx(blocks, style, title, base, out_path, template=None, post=None):
    """Собрать DOCX. Возвращает сводку {counts, warnings, out}. Ошибки входа — KitError (код 2), самопроверки — код 5."""
    tcfg = style.get("title") or {}
    mode = "template" if template else tcfg.get("mode", "generic")
    if mode not in ("template", "generic", "header", "none"):
        raise KitError(f"style.json: title.mode = {mode!r} не поддерживается", hint="допустимо: template, generic, header, none")
    template = template or (tcfg.get("template") if mode == "template" else None)
    if mode == "template" and not template:
        raise KitError("title.mode = template, но шаблон не указан", hint="укажи путь в style.json → title.template или параметром --template")
    metrics = page_metrics(style)
    warnings = []
    has_break, own_section = True, False
    if template:
        path = resolve_path(base, template)
        try:
            doc = Document(path)
        except Exception as e:
            raise KitError(f"шаблон не открывается как DOCX: {os.path.basename(path)} ({type(e).__name__})", hint="нужен файл .docx; .doc сначала пересохрани в Word")
        end, has_break, _ = title_mod.find_title_end(doc)
        title_mod.cut_body(doc, end)
        own_section = any(el.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None for el in list(doc.element.body)[:end])
        if not title_mod.apply_replacements(doc, end, title):
            warnings.append("титул шаблона не заполнен: в title.json нет replacements")
    else:
        doc = Document()
        body = doc.element.body
        for el in list(body):
            if el.tag != qn("w:sectPr"):
                body.remove(el)
        doc.core_properties.author = doc.core_properties.comments = doc.core_properties.title = ""
    upd = doc.settings.element.find(qn("w:updateFields"))
    if upd is not None:
        doc.settings.element.remove(upd)  # иначе Word при открытии спросит про обновление полей и зависнет в topdf
    ids = install_styles(doc, style, metrics[3])
    _setup_page(doc, metrics)
    if mode == "generic":
        if not title:
            warnings.append("title.json не найден или пуст — на титульном листе только название вуза")
        title_mod.build_generic(doc, title, metrics[3])
        has_break = False
    elif mode == "header":
        title_mod.build_header(doc, title)
    title_page = mode in ("template", "generic")
    body = doc.element.body
    mark = len([el for el in body if el.tag != qn("w:sectPr")])
    bld = _Builder(doc, style, base, ids, metrics)
    bld.warnings = warnings
    bld.no_break = not title_page
    bld.write(blocks)
    if title_page and not has_break and len(body) > mark and body[mark].tag != qn("w:sectPr"):
        first = body[mark]  # титул кончился без разрыва: тело начинается с новой страницы
        if first.tag == qn("w:p"):
            Paragraph(first, doc._body).paragraph_format.page_break_before = True
        else:
            first.addprevious(parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="{STYLE_BODY}"/></w:pPr><w:r><w:br w:type="page"/></w:r></w:p>'))
    on_title = (style.get("page_number") or {}).get("on_title", False)
    _page_number(doc, style, first_page_special=title_page and not own_section and not on_title)
    if post:
        _run_post(doc, base, post)
    finalize(doc, out_path)
    return {"counts": bld.counts, "warnings": bld.warnings, "out": out_path}
