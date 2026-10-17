# Титульный лист отчёта: граница титульной части шаблона курса, замены в ней по title.json, типовой титул МИСИС
# (generic), шапка строками (header) и печать абзацев титула для команды title.
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, Twips

from .common import KitError, clean_text

STYLE_TITLE = "RKTitle"  # стиль строк титула и шапки (в common.py его нет: проверке он не нужен)
BODY_STARTERS = ("СОДЕРЖАНИЕ", "ОГЛАВЛЕНИЕ", "ВВЕДЕНИЕ")
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
_W_P, _W_R, _W_T, _W_TAB, _W_BR = qn("w:p"), qn("w:r"), qn("w:t"), qn("w:tab"), qn("w:br")
_OFF = ("0", "false", "off")
HINT_TITLE = "посмотри абзацы титула командой: python report_kit.py title \"<шаблон.docx>\" — find должен совпадать с текстом буква в букву"


def _on(el):
    """Включён ли признак вида <w:pageBreakBefore/> (w:val="0" — выключен)."""
    return el is not None and el.get(qn("w:val")) not in _OFF


def _style_breaks(doc, style_id):
    """Стиль абзаца (или его предок по basedOn) начинает абзац с новой страницы?"""
    by_id = {s.get(qn("w:styleId")): s for s in doc.styles.element.findall(qn("w:style"))}
    seen = set()
    while style_id and style_id in by_id and style_id not in seen:
        seen.add(style_id)
        st = by_id[style_id]
        el = st.find(qn("w:pPr") + "/" + qn("w:pageBreakBefore"))
        if el is not None:
            return _on(el)
        based = st.find(qn("w:basedOn"))
        style_id = based.get(qn("w:val")) if based is not None else None
    return False


def _own(p, el):
    """Элемент принадлежит самому абзацу p, а не абзацу вложенной надписи."""
    return next(el.iterancestors(_W_P), None) is p


def _segments(p):
    """Куски текста абзаца по порядку: (элемент, текст). Табуляция — «\\t», перенос строки — «\\n»."""
    out = []
    for el in p.iter(_W_T, _W_TAB, _W_BR, qn("w:noBreakHyphen")):
        if el.getparent().tag != _W_R or not _own(p, el):
            continue  # w:tab внутри w:tabs — это позиция табуляции, не знак
        if el.tag == _W_T:
            out.append((el, el.text or ""))
        elif el.tag == _W_TAB:
            out.append((el, "\t"))
        elif el.tag == _W_BR:
            if el.get(qn("w:type")) in (None, "textWrapping"):
                out.append((el, "\n"))
        else:
            out.append((el, "-"))
    return out


def para_text(p):
    return "".join(t for _, t in _segments(p))


def find_title_end(doc):
    """Где кончается титульная часть: (число элементов тела в ней, есть ли разрыв страницы в конце, пояснение)."""
    body = doc.element.body
    for k, el in enumerate(body):
        if el.tag == qn("w:sectPr"):
            return k, False, "конец документа"
        if el.tag == qn("w:sdt") and k > 0 and el.find(".//" + qn("w:docPartGallery")) is not None:
            return k, False, "блок оглавления"
        if el.tag != _W_P:
            continue
        ppr = el.find(qn("w:pPr"))
        if any(b.get(qn("w:type")) == "page" and _own(el, b) for b in el.iter(_W_BR)):
            return k + 1, True, "разрыв страницы"
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            return k + 1, True, "разрыв раздела"
        if k == 0:
            continue
        text = para_text(el).strip()
        pbb = ppr.find(qn("w:pageBreakBefore")) if ppr is not None else None
        ps = ppr.find(qn("w:pStyle")) if ppr is not None else None
        if _on(pbb) or (pbb is None and ps is not None and _style_breaks(doc, ps.get(qn("w:val")))):
            return k, False, f"абзац «{text[:40]}» с новой страницы"
        if text.upper() in BODY_STARTERS:
            return k, False, f"абзац «{text[:40]}»"
    return len(body), False, "конец документа"


def cut_body(doc, end):
    """Удалить всё после титульной части (последний sectPr тела остаётся: в нём лист и колонтитулы)."""
    body = doc.element.body
    for el in list(body)[end:]:
        if el.tag != qn("w:sectPr"):
            body.remove(el)


def title_paragraphs(doc, end):
    """Абзацы титульной части по порядку (с абзацами таблиц и надписей) и их «тени» — копии надписей в mc:Fallback."""
    numbered, shadows = [], []
    for el in list(doc.element.body)[:end]:
        for p in el.iter(_W_P):
            in_fallback = any(a.tag == _MC + "Fallback" for a in p.iterancestors())
            (shadows if in_fallback else numbered).append(p)
    return numbered, shadows


def _replace_in(p, find, repl):
    """Заменить все вхождения find в абзаце: замена кладётся в первый задетый прогон (его формат сохраняется)."""
    count = 0
    start = 0
    while True:
        segs = _segments(p)
        pos = "".join(t for _, t in segs).find(find, start)
        if pos < 0:
            return count
        end, at, first = pos + len(find), 0, True
        for el, t in segs:
            a, b = at, at + len(t)
            at = b
            if b <= pos or a >= end:
                continue
            keep = t[:max(pos - a, 0)] + (repl if first else "") + t[max(min(end, b) - a, 0):]
            first = False
            if el.tag != _W_T:
                if not keep:
                    el.getparent().remove(el)
                    continue
                new = parse_xml(f"<w:t {nsdecls('w')}/>")  # замена началась со знака табуляции: вместо него — текст
                el.getparent().replace(el, new)
                el = new
            el.text = keep
            el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        count += 1
        start = pos + len(repl)


def apply_replacements(doc, end, title):
    """Замены title.json → replacements в титульной части. Возвращает число выполненных замен."""
    repls = (title or {}).get("replacements") or []
    if not isinstance(repls, list):
        raise KitError("title.json: replacements должен быть списком объектов {find, replace}")
    numbered, shadows = title_paragraphs(doc, end)
    for r in repls:
        if not isinstance(r, dict) or not isinstance(r.get("find"), str) or not r["find"] or "replace" not in r:
            raise KitError(f"title.json: замена должна быть объектом с непустым find и replace, а получено {r!r}")
        find, repl, par = r["find"], clean_text(r["replace"]), r.get("par")
        if par is not None and (not isinstance(par, int) or not 1 <= par <= len(numbered)):
            raise KitError(f"title.json: par={par!r} — в титульной части абзацы с 1 по {len(numbered)}", hint=HINT_TITLE)
        # команда title показывает табуляцию как \t и неразрывный пробел как ~ — принимаем find и в таком виде
        variants = [find]
        alt = find.replace("~", "\u00a0").replace("\\t", "\t")
        if alt != find:
            variants.append(alt)
        where = range(len(numbered)) if par is None else [par - 1]
        hits = []
        for f in variants:
            hits = [i for i in where if f in para_text(numbered[i])]
            if hits:
                find = f
                break
        if not hits:
            raise KitError(f"в титульной части шаблона не найден текст «{r['find']}»" + (f" (абзац {par})" if par else ""), hint=HINT_TITLE)
        if len(hits) > 1:
            nums = ", ".join(str(i + 1) for i in hits)
            raise KitError(f"текст «{r['find']}» есть в нескольких абзацах титула: {nums}", hint='добавь в замену номер абзаца: "par": N (номера печатает команда title)')
        _replace_in(numbered[hits[0]], find, repl)
        for s in shadows:
            if find in para_text(s):
                _replace_in(s, find, repl)
    return len(repls)


def _line(doc, text, bold=False, before=0, align="center", left=0):
    """Строка титула: стиль RKTitle, остальное — прямым оформлением."""
    p = doc.add_paragraph()
    p._p.get_or_add_pPr().style = STYLE_TITLE
    fmt = p.paragraph_format
    if align != "center":
        fmt.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if before:
        fmt.space_before = Pt(before)
    if left:
        fmt.left_indent = Twips(left)
    run = p.add_run(clean_text(text).replace("\n", " "))
    if bold:
        run.bold = True
    return p


def _val(title, key):
    v = (title or {}).get(key)
    return "" if v is None else clean_text(v).strip()


def _quoted(s):
    return s if s.startswith("«") or s.startswith('"') else f"«{s}»"


def _work_lines(title):
    """Строки «что за работа»: вид и номер, тема, дисциплина (пустые пропускаются)."""
    kind, num, theme, course = (_val(title, k) for k in ("work_kind", "work_number", "work_title", "course"))
    out = []
    head = " ".join(x for x in (kind.upper(), f"№ {num}" if num else "") if x)
    if head:
        out.append((head, True))
    if course:
        out.append((f"по дисциплине {_quoted(course)}", False))
    if theme:
        out.append((f"на тему: {_quoted(theme)}", False))
    return out


def _people_lines(title):
    student, group, variant, teacher = (_val(title, k) for k in ("student", "group", "variant", "teacher"))
    out = []
    if group:
        out.append(f"Выполнил: студент группы {group}")
        if student:
            out.append(student)
    elif student:
        out.append(f"Выполнил: {student}")
    if variant:
        out.append(f"Вариант {variant}")
    if teacher:
        out.append(f"Проверил: {teacher}")
    return out


def build_generic(doc, title, text_w):
    """Типовой титульный лист МИСИС. text_w — ширина текста в твипах (для блока справа и рамки нижней строки)."""
    _line(doc, "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ")
    _line(doc, "Федеральное государственное автономное образовательное учреждение высшего образования")
    _line(doc, "«Национальный исследовательский технологический университет «МИСИС»")
    _line(doc, "(НИТУ МИСИС)")
    first = True
    for key in ("institute", "department"):
        if _val(title, key):
            _line(doc, _val(title, key), before=18 if first else 0)
            first = False
    first = True
    for text, bold in _work_lines(title):
        _line(doc, text, bold=bold, before=110 if first else 6)
        first = False
    first = True
    for text in _people_lines(title):
        _line(doc, text, before=90 if first else 0, align="left", left=int(text_w * 0.5))
        first = False
    place = ", ".join(x for x in (_val(title, "city"), _val(title, "year")) if x)
    if place:
        # нижняя строка — рамкой, привязанной к низу поля: не зависит от того, сколько строк заняло остальное
        p = _line(doc, place)
        frame = parse_xml(f'<w:framePr {nsdecls("w")} w:w="{text_w}" w:hSpace="0" w:wrap="notBeside" w:vAnchor="margin" '
                          f'w:hAnchor="margin" w:xAlign="center" w:yAlign="bottom"/>')
        p._p.pPr.find(qn("w:pStyle")).addnext(frame)


def build_header(doc, title):
    """Шапка вместо титульного листа: те же сведения строками по центру, отчёт продолжается на той же странице."""
    lines = _work_lines(title) + [(text, False) for text in _people_lines(title)]
    last = None
    for text, bold in lines:
        last = _line(doc, text, bold=bold)
    if last is not None:
        last.paragraph_format.space_after = Pt(12)
    return len(lines)


def _show(text):
    return text.replace("\t", "\\t").replace("\u00a0", "~").replace("\n", "\\n")


def describe(template_path):
    """Строки для команды title: абзацы титульной части шаблона с номерами (их пишут в replacements → par)."""
    try:
        doc = Document(template_path)
    except Exception as e:
        raise KitError(f"шаблон не открывается как DOCX: {template_path} ({type(e).__name__})", hint="нужен файл .docx; .doc сначала пересохрани в Word")
    end, has_break, why = find_title_end(doc)
    numbered, _ = title_paragraphs(doc, end)
    lines = [f"Титульная часть: абзацев {len(numbered)}; граница — {why}."]
    for i, p in enumerate(numbered, 1):
        mark = ""
        if any(a.tag == qn("w:txbxContent") for a in p.iterancestors()):
            mark = " [надпись]"
        elif any(a.tag == qn("w:tc") for a in p.iterancestors()):
            mark = " [таблица]"
        lines.append(f"{i:>3}{mark}  {_show(para_text(p))}")
    lines.append('Замены пиши в title.json: "replacements": [{"find": "текст как выше", "replace": "новый текст", "par": N}] (par нужен, если текст есть в двух абзацах).')
    return lines
