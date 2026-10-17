# Разбор расписания экзаменационной сессии МИСИС (файлы Excel с misis.ru/students/session/) → JSON.
# Запуск: python session_xls.py <файл.xls|.xlsx>   (нужны xlrd для .xls и openpyxl для .xlsx)
#
# Вид файлов до первой публикации (декабрь 2026) неизвестен, поэтому разбираются два вероятных устройства листа:
#   1) сетка, как у расписания занятий: в строке-шапке названия групп по столбцам, ниже строки с датой (и временем)
#      и ячейками «Дисциплина / преподаватель / ауд.» под группой;
#   2) список: шапка «Дата | Время | Группа | Дисциплина | Преподаватель | Аудитория», по экзамену на строку
#      (в ячейке группы может быть несколько групп через запятую).
# Листы другого вида пропускаются с причиной — сервер скажет об этом владельцу, а студентам даст ссылку на страницу.
#
# Вывод: {"groups": {"<группа>": [{"date": "2027-01-12", "time": "10:00"|null, "name", "kind": "exam"|"consult"|"test",
#          "place", "teacher"}]}, "skipped": ["<лист>: причина"]}
import datetime as dt
import json
import re
import sys
import unicodedata

GROUP = re.compile(r"^[А-ЯЁA-Z]{2,8}-\d\d(?:-[\wА-ЯЁ]+)*$")
DATE = re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{2,4})")
TIME = re.compile(r"\b(\d{1,2})[:.](\d\d)\b")
PLACE = re.compile(r"(ауд\.?\s*[\wА-ЯЁ/-]+|\b[А-ЯЁ]-?\d{3,4}[а-яё]?\b|корп\.?\s*[\wА-ЯЁ]+)", re.I)
TEACHER = re.compile(r"[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?\s+[А-ЯЁ]\.\s?[А-ЯЁ]\.")
LAT = {"A": "А", "B": "В", "C": "С", "E": "Е", "H": "Н", "K": "К", "M": "М", "O": "О", "P": "Р", "T": "Т", "X": "Х", "Y": "У"}


def norm_group(s):
    s = re.sub(r"\s+", "", str(s or "")).upper()
    return "".join(LAT.get(ch, ch) for ch in s)


def text(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d.%m.%Y") + (v.strftime(" %H:%M") if isinstance(v, dt.datetime) and (v.hour or v.minute) else "")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return unicodedata.normalize("NFC", str(v)).strip()


def sheets(path):
    if path.lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for ws in wb.worksheets:
            yield ws.title, [[text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    else:
        import xlrd
        bk = xlrd.open_workbook(path)
        for sh in bk.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(text(xlrd.xldate_as_datetime(cell.value, bk.datemode)))
                    else:
                        row.append(text(cell.value))
                rows.append(row)
            yield sh.name, rows


def date_of(s):
    m = DATE.search(s)
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    y = y + 2000 if y < 100 else y
    try:
        return dt.date(y, mo, d).isoformat()
    except ValueError:
        return None


def time_of(s):
    m = TIME.search(DATE.sub(" ", s))  # «12.01.2027» — дата, а не время 12:01
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m and int(m.group(1)) < 24 and int(m.group(2)) < 60 else None


def kind_of(s):
    s = s.lower()
    return "consult" if "консульт" in s else "test" if "зач" in s and "экзам" not in s else "exam"


def entry(cell, date, time):
    """Ячейка сетки «Дисциплина / преподаватель / ауд.» → экзамен."""
    lines = [l.strip() for l in re.split(r"[\n;]+", cell) if l.strip()]
    if not lines:
        return None
    teacher = next((TEACHER.search(l).group(0) for l in lines if TEACHER.search(l)), "")
    place = next((PLACE.search(l).group(0) for l in lines if PLACE.search(l)), "")
    name = lines[0]
    for pat in (teacher, place):
        if pat:
            name = name.replace(pat, "")
    name = re.sub(r"\s*\((?:экзамен|зач[её]т[^)]*|консультация)\)\s*", " ", name, flags=re.I)
    name = re.sub(r"^(экзамен|консультация|зач[её]т)\s*[:—-]?\s*", "", name, flags=re.I).strip(" ,.-")
    if not name:
        return None
    return {"date": date, "time": time or time_of(cell), "name": name, "kind": kind_of(cell), "place": place, "teacher": teacher}


def grid(rows):
    """Сетка: шапка с группами по столбцам. Возвращает {группа: [экзамены]} или None, если шапки нет."""
    head = None
    for i, row in enumerate(rows[:30]):
        cols = {c: norm_group(v) for c, v in enumerate(row) if GROUP.match(norm_group(v))}
        if len(cols) >= 2:
            head, groups = i, cols
            break
    if head is None:
        return None
    # Пустые заголовки справа от группы (объединённые ячейки, подгруппы) относятся к ней же.
    width = max(len(r) for r in rows)
    owner, cur = {}, None
    for c in range(width):
        if c in groups:
            cur = groups[c]
        elif cur and head < len(rows) and c < len(rows[head]) and rows[head][c]:
            cur = None
        if cur:
            owner[c] = cur
    first_col = min(groups)
    out, date, time = {}, None, None
    for row in rows[head + 1:]:
        left = [v for v in row[:first_col]]
        d = next((date_of(v) for v in left if date_of(v)), None)
        t = next((time_of(v) for v in left if time_of(v)), None)
        if d:
            date, time = d, t
        elif t:
            time = t
        if not date:
            continue
        for c, g in owner.items():
            v = row[c] if c < len(row) else ""
            if not v or GROUP.match(norm_group(v)):
                continue
            e = entry(v, date, time)
            if e:
                out.setdefault(g, []).append(e)
    return out


def table(rows):
    """Список: шапка со словами «группа» и «дисциплина/предмет». Возвращает {группа: [экзамены]} или None."""
    cols = None
    for i, row in enumerate(rows[:30]):
        low = [v.lower() for v in row]
        if any("групп" in v for v in low) and any("дисципл" in v or "предмет" in v for v in low):
            find = lambda *keys: next((c for c, v in enumerate(low) if any(k in v for k in keys)), None)
            cols = {"group": find("групп"), "date": find("дата", "число"), "time": find("время", "начал"), "name": find("дисципл", "предмет"),
                    "teacher": find("преподават", "экзаменатор"), "place": find("ауд", "место"), "kind": find("вид", "форма")}
            head = i
            break
    if cols is None or cols["name"] is None:
        return None
    get = lambda row, k: (row[cols[k]] if cols[k] is not None and cols[k] < len(row) else "")
    out, date = {}, None
    for row in rows[head + 1:]:
        d = date_of(get(row, "date")) if cols["date"] is not None else next((date_of(v) for v in row if date_of(v)), None)
        if d:
            date = d
        name = get(row, "name")
        if not date or not name:
            continue
        gs = [norm_group(g) for g in re.split(r"[,;\n]+", get(row, "group")) if GROUP.match(norm_group(g))]
        kind_cell = get(row, "kind") or name
        e = {"date": date, "time": time_of(get(row, "time")) or time_of(" ".join(row)), "name": name.strip(), "kind": kind_of(kind_cell),
             "place": get(row, "place") or (PLACE.search(" ".join(row)).group(0) if PLACE.search(" ".join(row)) else ""), "teacher": get(row, "teacher")}
        for g in gs:
            out.setdefault(g, []).append(e)
    return out


def main(path):
    groups, skipped = {}, []
    for title, rows in sheets(path):
        rows = [[c for c in r] for r in rows if any(r)]
        got = grid(rows) or table(rows)
        if got is None:
            skipped.append(f"{title}: не нашёл ни групп по столбцам, ни шапки списка")
            continue
        if not got:
            skipped.append(f"{title}: группы есть, а строк с датами нет")
        for g, lst in got.items():
            seen = set()
            for e in lst:
                key = (e["date"], e["name"].lower())
                if key in seen:
                    continue
                seen.add(key)
                groups.setdefault(g, []).append(e)
    for lst in groups.values():
        lst.sort(key=lambda e: (e["date"], e["time"] or "99"))
    print(json.dumps({"groups": groups, "skipped": skipped}, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1])
