# Таблица преподавателя из курса Moodle (xlsx, xlsm, xls, csv) → JSON с ячейками как есть. Разбор (чья строка, какие
# столбцы — баллы, посещения, итог) — в course-sheets.mjs / sheets.mjs: там его проверяют тесты стенда.
# Запуск: python sheet_xls.py <файл>   (openpyxl — для .xlsx/.xlsm, xlrd — для .xls; defusedxml, если стоит, openpyxl
# подключает сам — защита от «XML-бомб» в чужих файлах).
#
# Вывод: {"sheets": [{"name", "hidden", "rows": [[ячейка]], "merged": [[r1, c1, r2, c2]] (с нуля, включительно),
#          "formulas": {"r:c": "=SUM(G17:AJ17)"}}], "truncated": bool}
#        или {"error": "причина"}.
# Ячейка: null | число | строка (NFC, без пробелов по краям) | {"d": "ГГГГ-ММ-ДД"} — дата.
# Формулы — только у ячеек, где Excel не сохранил значение (файл собран программой) или это сумма: итог студента часто
# посчитан формулой =SUM(…), и это главный признак столбца «Итого».
import csv
import datetime as dt
import io
import json
import re
import sys
import unicodedata

MAX_ROWS, MAX_COLS, MAX_SHEETS = 3000, 300, 40
SUMLIKE = re.compile(r"^=\s*(SUM|СУММ)\s*\(|^=\s*[A-Z]{1,3}\d+(\s*\+\s*[A-Z]{1,3}\d+){2,}\s*$", re.I)


def cell(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (dt.datetime, dt.date)):
        return {"d": v.strftime("%Y-%m-%d")}
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, (int, float)):
        if v != v or v in (float("inf"), float("-inf")):
            return None
        return int(v) if isinstance(v, float) and v.is_integer() else v
    s = unicodedata.normalize("NFC", str(v)).strip()
    return s or None


def trim(rows):
    rows = [r[:MAX_COLS] for r in rows[:MAX_ROWS]]
    out = []
    for r in rows:
        while r and r[-1] is None:
            r = r[:-1]
        out.append(r)
    while out and not out[-1]:
        out.pop()
    return out


def stored(ws):
    """Только ячейки, которые есть в файле (ws._cells): iter_rows создавал объект на каждую клетку до max_row×max_col,
    и файл в 8 КБ с одной ячейкой на (3000, 300) разбирался полторы минуты (проверка 16 окт)."""
    cells = getattr(ws, "_cells", None)
    if cells is None:  # другая версия openpyxl — как раньше, через iter_rows
        return [(c.row, c.column, c) for r in ws.iter_rows(max_row=min(ws.max_row or 0, MAX_ROWS), max_col=min(ws.max_column or 0, MAX_COLS)) for c in r]
    return [(r, c, x) for (r, c), x in cells.items() if r <= MAX_ROWS and c <= MAX_COLS]


def xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        wf = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        wf = None
    truncated = len(wb.worksheets) > MAX_SHEETS
    out = []
    for ws in wb.worksheets[:MAX_SHEETS]:
        if not hasattr(ws, "iter_rows") or getattr(ws, "max_row", 0) is None:
            continue  # листы-диаграммы
        truncated = truncated or (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLS
        grid = {}
        for r, c, x in stored(ws):
            v = cell(x.value)
            if v is not None:
                grid.setdefault(r - 1, {})[c - 1] = v
        rows = []
        for ri in range(max(grid.keys(), default=-1) + 1):
            row = grid.get(ri, {})
            rows.append([row.get(ci) for ci in range(max(row.keys(), default=-1) + 1)])
        rows = trim(rows)
        merged = []
        for m in ws.merged_cells.ranges:
            if m.min_row <= MAX_ROWS and m.min_col <= MAX_COLS:
                merged.append([m.min_row - 1, m.min_col - 1, min(m.max_row, MAX_ROWS) - 1, min(m.max_col, MAX_COLS) - 1])
        formulas = {}
        if wf is not None and ws.title in wf.sheetnames:
            for r, c, x in stored(wf[ws.title]):
                f = x.value
                if not isinstance(f, str) or not f.startswith("="):
                    continue
                ri, ci = r - 1, c - 1
                have = ri < len(rows) and ci < len(rows[ri]) and rows[ri][ci] is not None
                if not have or SUMLIKE.search(f):
                    formulas[f"{ri}:{ci}"] = f[:200]
        out.append({"name": ws.title, "hidden": ws.sheet_state != "visible", "rows": rows, "merged": merged, "formulas": formulas})
    return out, truncated


def xls(path):
    import xlrd
    try:
        bk = xlrd.open_workbook(path, formatting_info=True)
    except Exception:
        bk = xlrd.open_workbook(path)  # formatting_info не у всех файлов читается — тогда без объединённых ячеек
    truncated = bk.nsheets > MAX_SHEETS
    out = []
    for sh in bk.sheets()[:MAX_SHEETS]:
        truncated = truncated or sh.nrows > MAX_ROWS or sh.ncols > MAX_COLS
        rows = []
        for r in range(min(sh.nrows, MAX_ROWS)):
            row = []
            for c in range(min(sh.ncols, MAX_COLS)):
                x = sh.cell(r, c)
                if x.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(cell(xlrd.xldate_as_datetime(x.value, bk.datemode)))
                    except Exception:
                        row.append(cell(x.value))
                elif x.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    row.append(None)
                else:
                    row.append(cell(x.value))
            rows.append(row)
        merged = [[r1, c1, r2 - 1, c2 - 1] for (r1, r2, c1, c2) in getattr(sh, "merged_cells", []) if r1 < MAX_ROWS and c1 < MAX_COLS]
        out.append({"name": sh.name, "hidden": getattr(sh, "visibility", 0) != 0, "rows": trim(rows), "merged": merged, "formulas": {}})
    return out, truncated


NUM = re.compile(r"^[-+]?\d+(?:[.,]\d+)?$")
# «22.10», «08.10.2026» в CSV — даты занятий, а не числа 22,1 и 8,1 (проверка 16 окт); «1.50», «12,5» — числа.
DMY = re.compile(r"^(\d{1,2})\.(\d{2})(?:\.\d{2,4})?$")


def csv_value(v):
    m = DMY.match(v)
    if m and 1 <= int(m.group(1)) <= 31 and 1 <= int(m.group(2)) <= 12:
        return v
    return float(v.replace(",", ".")) if NUM.match(v) else v


def csvfile(path):
    raw = open(path, "rb").read(5 * 1024 * 1024)
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("кодировка CSV не распознана")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for r in csv.reader(io.StringIO(text), dialect):
        row = []
        for v in r[:MAX_COLS]:
            v = unicodedata.normalize("NFC", v).strip()
            row.append(None if not v else csv_value(v))
        rows.append([int(x) if isinstance(x, float) and x.is_integer() else x for x in row])
        if len(rows) >= MAX_ROWS:
            break
    return [{"name": "CSV", "hidden": False, "rows": trim(rows), "merged": [], "formulas": {}}], len(rows) >= MAX_ROWS


def main(path):
    low = path.lower()
    if low.endswith((".xlsx", ".xlsm")):
        sheets, truncated = xlsx(path)
    elif low.endswith(".xls"):
        sheets, truncated = xls(path)
    elif low.endswith(".csv"):
        sheets, truncated = csvfile(path)
    else:
        return {"error": "такой вид таблицы не поддерживается (нужен xlsx, xls или csv)"}
    return {"sheets": sheets, "truncated": truncated}


if __name__ == "__main__":
    try:
        res = main(sys.argv[1])
    except Exception as e:  # повреждённый или не тот файл — причина владельцу, без трассировки
        res = {"error": f"{type(e).__name__}: {str(e)[:200]}"}
    sys.stdout.write(json.dumps(res, ensure_ascii=False))
